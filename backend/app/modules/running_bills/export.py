from __future__ import annotations
from decimal import Decimal
from app.modules.running_bills.models import RunningBill
from io import BytesIO
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

def _money(value: Decimal, currency: str) -> str:
  return f"{currency} {value:,.2f}"

def build_running_bill_pdf(
  bill: RunningBill, organization_name: str, project_name: str, project_code: str | None, client_name: str | None,
) -> bytes:
  buffer = BytesIO()
  doc = SimpleDocTemplate(
    buffer, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm, leftMargin=15 * mm, rightMargin=15 * mm,
  )
  styles = getSampleStyleSheet()
  title_style = ParagraphStyle("BillTitle", parent=styles["Heading1"], fontSize=16, spaceAfter=2)
  small_style = ParagraphStyle("Small", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))

  elements = [
    Paragraph(organization_name, title_style),
    Paragraph(f"Running Bill (IPC) No. {bill.bill_number}", small_style),
    Spacer(1, 8),
  ]

  meta_rows = [
    ["Project", f"{project_name}{f' ({project_code})' if project_code else ''}"],
    ["Client", client_name or "—"],
    ["Period", f"{bill.period_start.isoformat()} to {bill.period_end.isoformat()}"],
    ["Status", bill.status.value],
  ]
  meta_table = Table(meta_rows, colWidths=[35 * mm, 130 * mm])
  meta_table.setStyle(TableStyle([
    ("FONTSIZE", (0, 0), (-1, -1), 9),
    ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#777777")),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
  ]))
  elements += [meta_table, Spacer(1, 12)]

  header = ["Description", "Unit", "Contract Qty", "Cum. %", "Cum. Qty", "This Period Qty", "Rate", "This Period Value"]
  rows = [header] + [
    [
      item.material_name, item.unit, f"{item.contract_quantity:,.2f}", f"{item.cumulative_percentage:.1f}%",
      f"{item.cumulative_quantity:,.2f}", f"{item.this_period_quantity:,.2f}",
      f"{item.unit_rate:,.2f}", f"{item.this_period_value:,.2f}",
    ]
    for item in bill.line_items
  ]

  items_table = Table(rows, colWidths=[45*mm, 14*mm, 20*mm, 14*mm, 20*mm, 22*mm, 18*mm, 25*mm], repeatRows=1)
  items_table.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
    ("FONTSIZE", (0, 0), (-1, -1), 8),
    ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dddddd")),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f4ee")]),
    ("TOPPADDING", (0, 0), (-1, -1), 4),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
  ]))
  elements += [items_table, Spacer(1, 14)]

  summary_rows = [
    ["Gross value this period", _money(bill.gross_value_this_period, bill.currency)],
    ["Gross value to date (cumulative)", _money(bill.gross_value_cumulative, bill.currency)],
    [f"Less: Retention ({bill.retention_percentage}%)", f"- {_money(bill.retention_this_period, bill.currency)}"],
  ]
  if bill.advance_recovery_amount:
    summary_rows.append(["Less: Advance recovery", f"- {_money(bill.advance_recovery_amount, bill.currency)}"])
  if bill.other_deductions_amount:
    summary_rows.append([bill.other_deductions_note or "Less: Other deductions", f"- {_money(bill.other_deductions_amount, bill.currency)}"])
  summary_rows.append(["NET PAYABLE THIS BILL", _money(bill.net_payable, bill.currency)])

  summary_table = Table(summary_rows, colWidths=[110 * mm, 48 * mm])
  summary_table.setStyle(TableStyle([
    ("FONTSIZE", (0, 0), (-1, -1), 9.5),
    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
    ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ("TOPPADDING", (0, 0), (-1, -1), 3),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
  ]))
  elements.append(summary_table)

  if bill.notes:
    elements += [Spacer(1, 14), Paragraph(f"<b>Notes:</b> {bill.notes}", small_style)]

  doc.build(elements)
  return buffer.getvalue()

def build_running_bill_xlsx(bill: RunningBill, organization_name: str, project_name: str, client_name: str | None) -> bytes:
  wb = Workbook()
  ws = wb.active
  ws.title = f"Bill {bill.bill_number}"

  bold = Font(bold=True)
  header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
  header_font = Font(bold=True, color="FFFFFF")

  ws.append([organization_name])
  ws["A1"].font = Font(bold=True, size=14)
  ws.append([f"Running Bill (IPC) No. {bill.bill_number}"])
  ws.append([f"Project: {project_name}"])
  ws.append([f"Client: {client_name or '—'}"])
  ws.append([f"Period: {bill.period_start} to {bill.period_end}"])
  ws.append([f"Status: {bill.status.value}"])
  ws.append([])

  header_row = ["Description", "Unit", "Contract Qty", "Cumulative %", "Cumulative Qty", "This Period Qty", "Rate", "This Period Value"]
  ws.append(header_row)
  for col_index in range(1, len(header_row) + 1):
    cell = ws.cell(row=ws.max_row, column=col_index)
    cell.fill = header_fill
    cell.font = header_font

  for item in bill.line_items:
    ws.append([
      item.material_name, item.unit, float(item.contract_quantity), float(item.cumulative_percentage),
      float(item.cumulative_quantity), float(item.this_period_quantity), float(item.unit_rate), float(item.this_period_value),
    ])

  ws.append([])
  summary_start = ws.max_row + 1
  ws.append(["Gross value this period", "", "", "", "", "", "", float(bill.gross_value_this_period)])
  ws.append(["Gross value to date", "", "", "", "", "", "", float(bill.gross_value_cumulative)])
  ws.append([f"Retention ({bill.retention_percentage}%)", "", "", "", "", "", "", -float(bill.retention_this_period)])
  if bill.advance_recovery_amount:
    ws.append(["Advance recovery", "", "", "", "", "", "", -float(bill.advance_recovery_amount)])
  if bill.other_deductions_amount:
    ws.append([bill.other_deductions_note or "Other deductions", "", "", "", "", "", "", -float(bill.other_deductions_amount)])
  ws.append(["NET PAYABLE", "", "", "", "", "", "", float(bill.net_payable)])
  for row in range(summary_start, ws.max_row + 1):
    ws.cell(row=row, column=1).font = bold
    ws.cell(row=row, column=8).font = bold

  for col_index, width in enumerate([32, 8, 14, 12, 14, 14, 10, 16], start=1):
    ws.column_dimensions[get_column_letter(col_index)].width = width

  buffer = BytesIO()
  wb.save(buffer)
  return buffer.getvalue()