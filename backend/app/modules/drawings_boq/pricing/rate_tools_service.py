from __future__ import annotations
import csv
import io
from datetime import date, datetime, timezone
from app.modules.drawings_boq.models import RateItem, RateBook
from sqlalchemy.ext.asyncio import AsyncSession
from app.modules.drawings_boq.pricing.rate_repository import RateRepository
from app.modules.audit.service import AuditLogService
from app.core.exceptions import TraceException
from app.modules.drawings_boq import boq_logic as logic
from uuid import UUID
from decimal import Decimal
from app.modules.drawings_boq.pricing.rate_service import RateBookService
from app.modules.drawings_boq.pricing.resolver import PriceResolver
from app.modules.drawings_boq.pricing.spec_match import stated_spec
from app.modules.drawings_boq.pricing.tools_schemas import RefreshRequest
from app.modules.drawings_boq.pricing import import_logic as il
from sqlalchemy import select
from app.modules.audit.models import AuditAction, AuditEntityType
from app.modules.drawings_boq.pricing.pricing_service import texts_for
from app.modules.drawings_boq.repository import BOQItemRepository, BOQVersionRepository

COMPARE_LINE_CAP = 2000
GAP_ITEM_CAP = 500
EXPORT_COLUMNS = ["work_item_code", "unit", "rate", "description", "trade", "specification", "csr_ref", "is_active"]

def _now():
  return datetime.now(timezone.utc)

def _d(item: RateItem) -> dict:
  return {"id": item.id, "work_item_code": item.work_item_code, "unit": item.unit, "rate": item.rate, "description": item.description, "trade": item.trade,
    "csr_ref": item.csr_ref, "analysis_id": item.analysis_id, "is_active": item.is_active, "specification": item.specification}

def _book_info(b: RateBook) -> dict:
  return {"id": str(b.id), "code": b.code, "name": b.name, "edition": b.edition, "immutable_version": b.immutable_version, "status": b.status,
    "effective_from": b.effective_from.isoformat() if b.effective_from else None, "scope": "ORGANIZATION" if b.organization_id else "SYSTEM"}

class RateToolsService:
  def __init__(self, session: AsyncSession):
    self.session = session
    self.repo = RateRepository(session)
    self.audit = AuditLogService(session)

  async def _book(self, org: UUID, book_id: UUID, *, for_update: bool = False) -> RateBook:
    book = await self.repo.get_book(book_id, org, for_update=for_update)
    if book is None:
      raise TraceException("Rate book not found.", status_code=404, code="RATE_BOOK_NOT_FOUND")
    return book

  async def _items(self, book_id: UUID) -> list[dict]:
    return [_d(i) for i in await self.repo.all_items(book_id)]

  async def export_csv(self, org: UUID, book_id: UUID) -> tuple[str, bytes]:
    book = await self._book(org, book_id)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(EXPORT_COLUMNS)
    for i in sorted(await self._items(book.id), key=lambda x: (x["trade"] or "", x["work_item_code"], x["unit"])):
      w.writerow([i["work_item_code"], i["unit"], logic.fmt(i["rate"], 2), i["description"] or "", i["trade"] or "", i["specification"] or "", i["csr_ref"] or "",
        "true" if i["is_active"] else "false"])
    return f"{book.code}_v{book.immutable_version}.csv", ("﻿" + out.getvalue()).encode("utf-8")

  async def export_xlsx(self, org: UUID, book_id: UUID) -> tuple[str, bytes]:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = await self._book(org, book_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(EXPORT_COLUMNS)
    for c in ws[1]:
      c.font = Font(bold=True)
    for i in sorted(await self._items(book.id), key=lambda x: (x["trade"] or "", x["work_item_code"], x["unit"])):
      ws.append([i["work_item_code"], i["unit"], float(i["rate"]), i["description"] or "", i["trade"] or "", i["specification"] or "", i["csr_ref"] or "", bool(i["is_active"])])
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGH", (22, 10, 14, 60, 22, 40, 16, 10)):
      ws.column_dimensions[col].width = width
    info = wb.create_sheet("Book")
    
    for k, v in (("Code", book.code), ("Name", book.name), ("Version", book.immutable_version), ("Edition", book.edition), ("Status", book.status),
      ("Currency", book.currency), ("Jurisdiction", book.jurisdiction), ("Province", book.province), ("City", book.city),
      ("Effective from", book.effective_from.isoformat() if book.effective_from else None), ("Effective to", book.effective_to.isoformat() if book.effective_to else None),
      ("Content hash", book.content_hash), ("Exported", _now().isoformat(timespec="seconds"))):
      info.append([k, v])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 48
    buf = io.BytesIO()
    wb.save(buf)
    return f"{book.code}_v{book.immutable_version}.xlsx", buf.getvalue()

  async def compare(self, org: UUID, a_id: UUID, b_id: UUID, match_on: str) -> dict:
    a, b = await self._book(org, a_id), await self._book(org, b_id)
    out = il.compare_books(await self._items(a.id), await self._items(b.id), match_on=match_on)
    lines = out["lines"]
    return {"book_a": _book_info(a), "book_b": _book_info(b), "match_on": match_on, "counts": out["counts"], "matched": out["matched"],
      "median_change_pct": out["median_change_pct"], "overall_ratio": out["overall_ratio"], "suggested_factors": out["suggested_factors"],
      "lines": lines[:COMPARE_LINE_CAP], "truncated": len(lines) > COMPARE_LINE_CAP}

  async def history(self, org: UUID, book_id: UUID, work_item_code: str, unit: str | None) -> dict:
    book = await self._book(org, book_id)
    q = select(RateBook).where(RateBook.code == book.code, RateBook.status.in_(("ACTIVE", "SUPERSEDED", "ARCHIVED")))
    q = q.where(RateBook.organization_id == book.organization_id) if book.organization_id else q.where(RateBook.organization_id.is_(None))
    books = list((await self.session.execute(q)).scalars().all())
    if book.status == "DRAFT":
      books.append(book)
    by_id = {b.id: b for b in books}
    want = logic.normalise_unit(unit) if unit else None
    entries: list[dict] = []
    
    if by_id:
      rows = (await self.session.execute(select(RateItem).where(RateItem.rate_book_id.in_(list(by_id)), RateItem.work_item_code == work_item_code))).scalars().all()
      for it in rows:
        if want and logic.normalise_unit(it.unit) != want:
          continue
        b = by_id[it.rate_book_id]
        entries.append({"book_id": str(b.id), "book_code": b.code, "version": b.immutable_version, "edition": b.edition, "status": b.status,
          "effective_from": b.effective_from.isoformat() if b.effective_from else None, "unit": it.unit, "rate": str(it.rate)})
    return {"work_item_code": work_item_code, "unit": unit, "points": il.history_series(entries)}

  async def refresh(self, org: UUID, user_id: UUID, book_id: UUID, payload: RefreshRequest) -> dict:
    target = await self._book(org, book_id, for_update=True)
    if target.organization_id != org:
      raise TraceException("System rate books are read-only.", status_code=403, code="RATE_BOOK_READ_ONLY")
    if target.status != "DRAFT":
      raise TraceException("Only a DRAFT rate book can be refreshed. Start a new version of it first.", status_code=409, code="RATE_BOOK_IMMUTABLE")
    source = await self._book(org, payload.source_book_id)
    if source.id == target.id:
      raise TraceException("Choose a different edition to refresh from.", status_code=422, code="RATE_REFRESH_SAME_BOOK")
    if source.status == "DRAFT":
      raise TraceException("Refresh from a published edition, not from a draft.", status_code=422, code="RATE_REFRESH_SOURCE_DRAFT")
    only = {t for t in (payload.trades or [])} or None
    rows = await self.repo.all_items(target.id)
    by_id = {r.id: r for r in rows}
    
    plan = il.plan_refresh([_d(r) for r in rows], await self._items(source.id), match_on=payload.match_on, only_trades=only)
    analyses_out = {"recomputed": 0, "changed": [], "failed": []}
    sp = await self.session.begin_nested()
    try:
      stamp = {"book_id": str(source.id), "code": source.code, "version": source.immutable_version, "edition": source.edition, "at": _now().isoformat(timespec="seconds")}
      for ch in plan["changes"]:
        it = by_id[ch["id"]]
        it.rate = logic.q2(Decimal(ch["new"]))
        it.extra = {**(it.extra or {}), "refreshed_from": stamp, "previous_rate": ch["old"]}
        
      await self.session.flush()
      if payload.recompute_analyses:
        svc = RateBookService(self.session)
        analyses = await self.repo.list_analyses(target.id)
        for a in analyses:
          try:
            res = await svc.breakdown(org, a)
          except TraceException as exc:
            analyses_out["failed"].append({"code": a.code, "error": exc.message if hasattr(exc, "message") else str(exc)})
            continue
          old = a.computed_rate
          a.computed_rate, a.computed_at = res["rate_value"], _now()
          analyses_out["recomputed"] += 1
          linked = next((r for r in rows if r.analysis_id == a.id), None)
          if linked is not None and linked.rate != a.computed_rate:
            analyses_out["changed"].append({"code": a.code, "unit": a.unit, "old": str(linked.rate), "new": str(a.computed_rate)})
            linked.rate = logic.q2(a.computed_rate)
          elif old != a.computed_rate:
            analyses_out["changed"].append({"code": a.code, "unit": a.unit, "old": None if old is None else str(old), "new": str(a.computed_rate)})
        await self.session.flush()
        
    except BaseException:
      await sp.rollback()
      raise
    if payload.apply:
      await sp.commit()
      await self.session.commit()
      await self.audit.log(org, user_id, AuditEntityType.RATE_BOOK, target.id, AuditAction.UPDATE,
        f"Refreshed {len(plan['changes'])} rate(s) of {target.code} v{target.immutable_version} from {source.code} v{source.immutable_version}")
    else:
      await sp.rollback()
    analyses_out["changed"] = analyses_out["changed"][:200]
    return {"applied": payload.apply, "source": _book_info(source), "match_on": payload.match_on, "updated": len(plan["changes"]), "unchanged": plan["unchanged"],
      "unmatched": len(plan["unmatched"]), "skipped": len(plan["skipped"]), "changes": [{**c, "id": str(c["id"])} for c in plan["changes"][:2000]],
      "unmatched_sample": [{**u, "id": str(u["id"])} for u in plan["unmatched"][:100]], "analyses": analyses_out}

  async def rate_gaps(self, org: UUID, version_id: UUID, *, as_of: date | None, rate_book_ids: list[UUID] | None) -> dict:
    version = await BOQVersionRepository(self.session).get_by_id_and_org(version_id, org)
    if version is None:
      raise TraceException("BOQ version not found.", status_code=404, code="BOQ_VERSION_NOT_FOUND")
    items = await BOQItemRepository(self.session).list_by_version(version.id, org)
    on = as_of or date.today()
    resolver = await PriceResolver.load(self.session, org, version.project_id, as_of=on, codes={i.work_item_code for i in items if i.work_item_code}, rate_book_ids=rate_book_ids)
    groups: dict[tuple, dict] = {}
    priced = unpriced = resolvable = 0
    
    for item in items:
      if item.unit_rate is not None:
        priced += 1
        continue
      unpriced += 1
      res = resolver.resolve(item.work_item_code, texts_for(item), item.unit, spec=stated_spec(item))
      if res.found:
        resolvable += 1
        continue
    
      code = (item.work_item_code or "").strip() or None
      reason = "NO_CODE" if code is None else "UNIT_MISMATCH" if res.unit_mismatch else "NO_RATE"
      key = (code, logic.normalise_unit(item.unit), reason)
      g = groups.setdefault(key, {"work_item_code": code, "unit": key[1], "reason": reason, "lines": 0, "quantity": Decimal("0"),
        "sample_description": item.description or item.material_name, "available_units": set(), "available_specs": set(), "stated": stated_spec(item), "item_ids": []})
      g["lines"] += 1
      g["quantity"] += Decimal(item.quantity or 0)
      if len(g["item_ids"]) < GAP_ITEM_CAP:
        g["item_ids"].append(item.id)
      if code:
        for b in resolver.books:
          for e in resolver.items.get((b.id, code), []):
            g["available_units"].add(e.unit)
          for _, composed in resolver.spec_index.get((b.id, code), []):
            g["available_specs"].add(composed)
        for o in resolver.overrides.get(code, []):
          g["available_units"].add(o.unit)
          
    gaps = sorted(groups.values(), key=lambda g: (-g["lines"], g["work_item_code"] or ""))
    for g in gaps:
      g["quantity"] = logic.fmt(g["quantity"], 3) if hasattr(logic, "fmt") else str(g["quantity"])
      g["available_units"] = sorted(g["available_units"])
      g["available_specs"] = sorted(g["available_specs"])[:12]
    return {"boq_version_id": version.id, "as_of": on.isoformat(), "rate_books": resolver.stack_summary(), "priced": priced, "unpriced": unpriced,
      "resolvable": resolvable, "gaps": gaps}
