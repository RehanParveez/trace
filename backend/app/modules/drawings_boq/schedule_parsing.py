from __future__ import annotations
from decimal import Decimal, InvalidOperation
import re
from app.modules.drawings_boq.boq_logic import q6, UNIT_TABLE
import csv
import io

MAX_ROWS = 1000
SURFACES = ("FLOOR", "WALL", "CEILING", "SKIRTING", "DADO")
SURFACE_UNIT = {"FLOOR": "m2", "WALL": "m2", "CEILING": "m2", "DADO": "m2", "SKIRTING": "m"}
KIND_DEFAULT_WORK_ITEM = {"DOOR": "DOR-NOS", "WINDOW": "WIN-NOS"}
CSV_CONFIDENCE = Decimal("0.9")
CONFIDENCE_WORDS = {"high": Decimal("0.85"), "medium": Decimal("0.6"), "low": Decimal("0.35")}

class ScheduleParseError(Exception):
  pass

def norm_key(text) -> str:
  return re.sub(r"[^a-z0-9]", "", str(text or "").lower())

norm_mark = norm_key

def _clean(value) -> str | None:
  if value is None:
    return None
  text = str(value).strip()
  return text or None

_UNIT_ALIAS = {
  "m3": "m3", "cum": "m3", "cbm": "m3", "cft": "cft", "cuft": "cft", "ft3": "cft",
  "m2": "m2", "sqm": "m2", "sft": "sft", "sqft": "sft", "ft2": "sft",
  "m": "m", "rm": "m", "rmt": "m", "lm": "m", "mtr": "m", "rft": "rft", "rf": "rft", "ft": "rft",
  "kg": "kg", "kgs": "kg",
  "nos": "nos", "no": "nos", "each": "nos", "ea": "nos", "pcs": "nos", "pc": "nos", "unit": "nos", "units": "nos",
}
_TONNE = frozenset({"ton", "tons", "tonne", "tonnes", "mt"})

def canonicalize(unit_text, quantity) -> tuple[str | None, Decimal | None, str | None]:
  if not unit_text or quantity is None:
    return None, None, None
  key = re.sub(r"[\s\.]+", "", str(unit_text).lower()).replace("²", "2").replace("³", "3")
  if key in _TONNE:
    return "kg", q6(Decimal(quantity) * 1000), "Tonne assumed to be 1000 kg."
  name = _UNIT_ALIAS.get(key)
  if name is None:
    return None, None, None
  canonical, factor = UNIT_TABLE[name]
  return canonical, q6(Decimal(quantity) / factor), None

def parse_decimal(value) -> Decimal | None:
  if value is None or isinstance(value, bool):
    return None
  text = re.sub(r"[^\d.\-]", "", str(value).replace(",", ""))
  if not text or text in {"-", "."}:
    return None
  try:
    number = Decimal(text)
  except InvalidOperation:
    return None
  return number if number.is_finite() else None

_FT_IN = re.compile(r"(\d+(?:\.\d+)?)\s*(?:'|ft|feet)\s*-?\s*(?:(\d+(?:\.\d+)?)\s*(?:\"|in|inch|inches)?)?")
_Q3 = Decimal("0.001")

def _length_mm(token: str, hint: str | None) -> tuple[Decimal | None, str | None]:
  t = token.strip().lower().replace("’", "'").replace("”", '"')
  m = _FT_IN.fullmatch(t)
  if m:
    return (Decimal(m.group(1)) * Decimal("304.8") + Decimal(m.group(2) or 0) * Decimal("25.4")).quantize(_Q3), None
  m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(mm|cm|m)?", t)
  if not m:
    return None, None
  n, unit = Decimal(m.group(1)), (m.group(2) or hint)
  if n <= 0:
    return None, None
  if unit == "mm":
    return n.quantize(_Q3), None
  if unit == "cm":
    return (n * 10).quantize(_Q3), None
  if unit == "m":
    return (n * 1000).quantize(_Q3), None
  if n >= 100:
    return n.quantize(_Q3), None
  if n < 10 and n != n.to_integral_value():
    return (n * 1000).quantize(_Q3), "Size read as metres."
  if n <= 20:
    return (n * Decimal("304.8")).quantize(_Q3), "Size read as feet."
  return n.quantize(_Q3), None

def parse_size(text) -> tuple[Decimal | None, Decimal | None, str | None]:
  if not text:
    return None, None, None
  s = str(text).lower().replace("×", "x").replace("*", "x")
  parts = [p for p in re.split(r"\s*x\s*", s) if p.strip()]
  if len(parts) != 2:
    return None, None, None
  hint = re.search(r"(mm|cm|m)\s*$", s)
  hint = hint.group(1) if hint else None
  w, wn = _length_mm(parts[0], hint)
  h, hn = _length_mm(parts[1], hint)
  return w, h, wn or hn

def infer_surface(text) -> str | None:
  t = str(text or "").lower()
  for keyword, surface in (("skirt", "SKIRTING"), ("dado", "DADO"), ("ceil", "CEILING"), ("floor", "FLOOR"), ("wall", "WALL")):
    if keyword in t:
      return surface
  return None

def suggest_finish_work_item(surface: str | None, text: str | None) -> str | None:
  t = (text or "").lower()
  if surface == "FLOOR":
    return "FIN-FLOOR"
  if surface == "SKIRTING":
    return "FIN-SKIRT"
  if surface == "DADO":
    return "FIN-DADO"
  paint = any(k in t for k in ("paint", "emulsion", "distemper"))
  if surface == "CEILING":
    return "FIN-CEIL-PAINT" if paint else "FIN-CEIL-PLASTER"
  if surface == "WALL":
    return "FIN-PAINT" if paint else ("FIN-PLASTER" if "plaster" in t else None)
  return None

_HEADERS = {

  "mark": {"mark", "type mark", "tag", "ref", "reference", "id", "identifier", "door mark", "window mark", "door no", "window no", "door number", "window number", "door type", "window type",
    "type mark", "schedule mark", "schedule ref", "schedule reference"},
  
  "description": {"description", "desc", "item", "item description", "item desc", "name", "item name", "particulars", "particular", "details", "detail", "description of work", 
    "work description", "scope", "scope of work"},

  "location": {"location", "room", "room name", "space", "space name", "zone", "zone name", "area", "area name", "room no", "room number", "space no", "space number", "location name", "place"},
  
  "level": {"level", "storey", "storey level", "story", "story level", "floor", "floor level", "floor no", "floor number", "level no", "level number", "storey no", 
    "storey number", "elevation", "building level"},

  "unit": {"unit", "units", "uom", "unit of measure", "unit of measurement", "measurement unit", "measure", "measuring unit"},
  "quantity": {"quantity", "qty", "qnty", "quant", "nos", "no", "count", "no of", "no. of", "number of", "number", "total", "total quantity", "total qty", "required quantity", "required qty"},
  "width": {"width", "w", "width mm", "w mm", "width (mm)", "width (m)", "w (mm)", "w (m)", "clear width", "opening width", "overall width"},
  "height": {"height", "h", "height mm", "h mm", "ht", "ht.", "height (mm)", "height (m)", "h (mm)", "h (m)", "clear height", "opening height", "overall height"},

  "size": {"size", "size mm", "size (mm)", "size (m)", "dimensions", "dimension", "dim", "dims", "w x h", "wxh", "w × h", "width x height", "width × height",
    "opening size", "opening dimensions"},

  "work_item": {"work item", "work item code", "item code", "item no", "item number", "wi code", "work code", "work item no", "work item number", "boq code", "boq item", "boq item code",
    "bill item", "bill item code", "wbs", "wbs code", "cost code", "activity code"},

  "surface": {"surface", "finish surface", "surface type", "surface name", "applies to", "applied to", "application surface", "substrate", "base surface", "wall surface",
    "floor surface", "ceiling surface"},

  "finish": {"finish", "finish type", "material", "finish name"},
    "dado_height": {"dado height", "dado ht", "dado h"},

}
_HEADER_LOOKUP = {name: key for key, names in _HEADERS.items() for name in names}

_SURFACE_HEADERS = {
  "floor": "FLOOR", "floors": "FLOOR", "flooring": "FLOOR", "floor finish": "FLOOR",
  "wall": "WALL", "walls": "WALL", "wall finish": "WALL", "walls finish": "WALL",
  "ceiling": "CEILING", "ceilings": "CEILING", "ceiling finish": "CEILING",
  "skirting": "SKIRTING", "skirting finish": "SKIRTING",
  "dado": "DADO", "dado finish": "DADO",
}
_EMPTY_CELLS = {"-", "--", "n/a", "na", "nil", "none", "null"}

def _split_finishes(value) -> list[str]:
  text = str(value or "").strip()
  if not text or text.lower() in _EMPTY_CELLS:
    return []
  return [p.strip() for p in re.split(r"\s*(?:\+|&|\band\b)\s*", text, flags=re.IGNORECASE) if p.strip()]

def _header_key(cell) -> str:
  return re.sub(r"[^a-z0-9]+", " ", str(cell or "").lower()).strip()

def read_table(content: bytes, filename: str) -> list[list]:
  name = (filename or "").lower()
  if name.endswith(".xlsx"):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
  try:
    text = content.decode("utf-8-sig")
  except UnicodeDecodeError:
    text = content.decode("latin-1")
  try:
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
  except csv.Error:
    dialect = csv.excel
  return [list(r) for r in csv.reader(io.StringIO(text), dialect)]

def _single_length(value) -> tuple[Decimal | None, str | None]:
  return _length_mm(str(value), None) if value not in (None, "") else (None, None)

def build_row(cells: dict, kind: str) -> dict | None:
  mark = _clean(cells.get("mark"))
  finish_name = _clean(cells.get("finish"))
  description = _clean(cells.get("description")) or finish_name
  location = _clean(cells.get("location"))
  
  if not (mark or description or location):
    return None
  qty = parse_decimal(cells.get("quantity"))
  if qty is not None and qty < 0:
    qty = None
    
  notes: list[str] = []
  width = height = None
  if cells.get("size"):
    width, height, note = parse_size(cells["size"])
    if note:
      notes.append(note)
      
  for key in ("width", "height"):
    value, note = _single_length(cells.get(key))
    if value is not None:
      width, height = (value, height) if key == "width" else (width, value)
      if note:
        notes.append(note)
        
  surface = None
  if kind == "FINISH":
    surface = infer_surface(cells.get("surface")) or infer_surface(description) or infer_surface(finish_name)
  return {
    "mark": mark, "description": description, "location_text": location, "level_text": _clean(cells.get("level")),
    "unit": _clean(cells.get("unit")), "quantity": qty, "width_mm": width, "height_mm": height,
    "surface": surface, "finish_name": finish_name, "work_item_code": _clean(cells.get("work_item")),
    "notes": notes, "confidence": CSV_CONFIDENCE,
  }

def rows_from_table(table: list[list], kind: str) -> list[dict]:
  header_idx, mapping, wide = None, {}, {}
  for i, row in enumerate(table[:10]):
    found = {j: _HEADER_LOOKUP[_header_key(c)] for j, c in enumerate(row) if _header_key(c) in _HEADER_LOOKUP}
    cols = {j: _SURFACE_HEADERS[_header_key(c)] for j, c in enumerate(row) if _header_key(c) in _SURFACE_HEADERS} \
      if kind == "FINISH" and "surface" not in found.values() else {}
    anchored = bool(set(found.values()) & {"mark", "description", "location", "surface"})
    if (len(found) >= 2 and anchored) or (cols and "location" in found.values()):
      header_idx, mapping, wide = i, found, cols
      break
    
  if header_idx is None:
    raise ScheduleParseError("Could not find a header row. Expected columns such as mark, description, unit, quantity.")
  out: list[dict] = []
  for raw in table[header_idx + 1:]:
    cells: dict = {}
    for col, key in mapping.items():
      if col < len(raw) and raw[col] not in (None, ""):
        cells.setdefault(key, raw[col])
    raw_text = " | ".join(str(c) for c in raw if c not in (None, ""))[:2000]
    built: list[dict] = []
    
    if wide:
      base = {k: v for k, v in cells.items() if k not in ("dado_height", "height", "width", "size")}
      for col, surface in wide.items():
        if col >= len(raw):
          continue
        for name in _split_finishes(raw[col]):
          c = {**base, "surface": surface, "finish": name}
          if surface == "DADO" and cells.get("dado_height"):
            c["height"] = cells["dado_height"]
          row = build_row(c, kind)
          if row is not None:
            built.append(row)
            
    else:
      row = build_row(cells, kind)
      if row is not None:
        built.append(row)
    for row in built:
      row["raw_text"] = raw_text
      out.append(row)
    if len(out) > MAX_ROWS:
      raise ScheduleParseError(f"The schedule has more than {MAX_ROWS} rows. Split it and import in parts.")
  if not out:
    raise ScheduleParseError("No data rows found under the header.")
  return out

def build_schedule_rows_prompt(kind: str, text: str) -> str:
  extra = (
    ' For finish schedules emit ONE item per room and surface, with "surface" one of '
    'FLOOR, WALL, CEILING, SKIRTING, DADO and "finish" the finish as written.'
    if kind == "FINISH" else ""
  )
  return (
    f"You are reading raw text extracted from a construction drawing PDF. Find the {kind} schedule table "
    "and report its rows exactly as written. Do NOT estimate, infer or calculate any value that is not literally "
    "in the text, and never convert units: copy sizes into size_text as written. If there is no such schedule, "
    "return an empty items array." + extra + "\n\nRespond with strict JSON only, no markdown:\n"
    '{"items": [{"mark": "..." or null, "description": "..." or null, "location": "..." or null, '
    '"unit": "..." or null, "quantity": <number or null>, "size_text": "..." or null, '
    '"surface": "..." or null, "finish": "..." or null, "confidence": "high" | "medium" | "low"}]}\n\n'
    f"Extracted drawing text:\n{text}"
  )

def parse_ai_rows(raw_output: dict | None, kind: str) -> list[dict]:
  items = (raw_output or {}).get("items")
  if not isinstance(items, list):
    return []
  out: list[dict] = []
  for r in items[:MAX_ROWS]:
    if not isinstance(r, dict):
      continue
    cells = {"mark": r.get("mark"), "description": r.get("description"), "location": r.get("location"),
      "unit": r.get("unit"), "quantity": r.get("quantity"), "size": r.get("size_text"),
      "surface": r.get("surface"), "finish": r.get("finish")}
    row = build_row({k: v for k, v in cells.items() if v not in (None, "")}, kind)
    if row is None:
      continue
    row["confidence"] = CONFIDENCE_WORDS.get(str(r.get("confidence") or "").lower(), CONFIDENCE_WORDS["low"])
    row["raw_text"] = str({k: v for k, v in r.items() if v is not None})[:2000]
    out.append(row)
  return out

def table_from_layout_text(text: str, header_probe=None) -> list[list]:
  lines = [ln.rstrip() for ln in (text or "").splitlines() if ln.strip()]

  def cells_of(line):
    return [(m.start(), m.end(), m.group(0).strip()) for m in re.finditer(r"\S+(?: \S+)*", line)]

  def probe(line):
    row = [c[2] for c in cells_of(line)]
    if header_probe is not None:
      return header_probe(row)
    found = {_HEADER_LOOKUP[_header_key(c)] for c in row if _header_key(c) in _HEADER_LOOKUP}
    return len(found) >= 2 and bool(found & {"mark", "description", "location", "surface"})

  start = next((i for i, ln in enumerate(lines) if len(cells_of(ln)) >= 2 and probe(ln)), None)
  if start is None:
    raise ScheduleParseError("No schedule table header was found in the PDF text.")
  
  head = cells_of(lines[start])
  centres = [(a + b) / 2 for a, b, _ in head]
  table = [[c[2] for c in head]]
  for ln in lines[start + 1:]:
    toks = cells_of(ln)
    if len(toks) < 2:
      continue
    row = [""] * len(head)
    for a, b, t in toks:
      j = min(range(len(centres)), key=lambda k: abs(centres[k] - (a + b) / 2))
      row[j] = (row[j] + " " + t).strip()
    table.append(row)
  return table

def rows_from_pdf_text(text: str, kind: str) -> list[dict]:
  rows = rows_from_table(table_from_layout_text(text), kind)
  for r in rows:
    r["confidence"] = Decimal("0.7")
    r["notes"] = [*r.get("notes", []), "Read from PDF text; check column alignment."]
  return rows