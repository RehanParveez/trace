from __future__ import annotations
import os
import yaml
import re
from app.modules.drawings_boq.lab.bill import BillLine, parse_bill, _text, _number
from app.modules.drawings_boq.lab import tolerances as T
from app.modules.drawings_boq.takeoff.proof import case as CASE
from app.modules.drawings_boq.lab.aliases import bill_key

CASE_CLASSES = ("DRAWING", "QS_ERROR", "CONVENTION", "NEEDS_DECISION")
FORMATS = ("WORKSHEET_MAP", "SUMMARY")
FACT_KINDS = ("COUNT_BY_MARK", "COUNT_BY_LABEL", "MEMBER_SIZE", "LEVEL", "ROOM_SIZE", "AREA", "SLAB_THICKNESS", "BEAM_SIZE", "TANK_SIZE", "ELEC_COUNT",
  "WALL_RUN", "FOOTING_SIZE", "COLUMN_SIZE", "STEP_TABLE")
CASES_DIR_ENV = "TRACE_CASES_DIR"
HERE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cases")

def _yaml(path: str):
  with open(path, encoding="utf-8") as fh:
    return yaml.safe_load(fh)

def load_case(folder: str) -> dict:
  path = os.path.join(folder, "case.yaml")
  if not os.path.isfile(path):
    raise ValueError(f"{folder} has no case.yaml.")
  try:
    raw = _yaml(path) or {}
  except yaml.YAMLError as e:
    raise ValueError(f"{path} is not valid YAML: {e}.")
  if not raw.get("name"):
    raise ValueError(f"{path} has no name.")
  pdfs = raw.get("pdfs") or []

  if not pdfs or not all(isinstance(p, str) and p for p in pdfs):
    raise ValueError(f"{path} lists no PDF.")
  wb = raw.get("workbook") or {}
  if wb and wb.get("format") not in FORMATS:
    raise ValueError(f"{path}: workbook format must be one of {', '.join(FORMATS)}.")
  out = {"name": str(raw["name"]), "tier": str(raw.get("tier") or "C"), "pdfs": [str(p) for p in pdfs], "sha256": dict(raw.get("sha256") or {}), "workbook": wb or None,
    "block_floor": dict(raw.get("block_floor") or {}), "company_rules": raw.get("company_rules"), "note": raw.get("note") or "", "folder": folder}
  out["expected"] = load_expected(os.path.join(folder, "expected.yaml"))
  out["facts"] = load_facts(os.path.join(folder, "facts.yaml"))
  return out

def load_company_rules(case: dict) -> tuple:
  name = case.get("company_rules")
  if not name:
    return [], [], None
  
  from app.engine.measure import rulebook as RB
  path = os.path.join(case["folder"], str(name))
  if not os.path.isfile(path):
    return [], [f"{name} is not in the case folder."], str(name)
  rows, problems = RB.load_rules(path)
  return rows, problems, str(name)

def load_expected(path: str) -> list:
  if not os.path.isfile(path):
    return []

  doc = _yaml(path) or []
  if not isinstance(doc, list):
    raise ValueError(f"{path} must be a list of bill lines.")
  out, seen = [], set()
  for i, e in enumerate(doc):
    if not isinstance(e, dict) or not e.get("key"):
      raise ValueError(f"{path}, entry {i + 1}: every line needs a key.")
    key = str(e["key"])
    if key in seen:
      raise ValueError(f"{path}: the key {key} is used twice.")
    seen.add(key)
    cls = str(e.get("class") or "DRAWING").upper()
    if cls not in CASE_CLASSES:
      raise ValueError(f"{path}, {key}: class must be one of {', '.join(CASE_CLASSES)}.")

    qs = e.get("qs")
    drawing = e.get("drawing")
    for name, v in (("qs", qs), ("drawing", drawing)):
      if v is not None and (isinstance(v, bool) or not isinstance(v, (int, float))):
        raise ValueError(f"{path}, {key}: {name} must be a number or null.")
    tol = e.get("tolerance_pct")
    out.append({"key": key, "item": str(e.get("item") or ""), "unit": str(e.get("unit") or ""), "qs": None if qs is None else float(qs), "drawing": None if drawing is None else float(drawing),
      "class": cls, "tolerance_pct": None if tol is None else float(tol), "note": str(e.get("note") or ""), "family": e.get("family"), "families": [str(f) for f in (e.get("families") or [])], "floor": e.get("floor"), "grade": e.get("grade"),
      "match": [str(w).lower() for w in (e.get("match") or [])]})
  return out

def load_facts(path: str) -> list:
  if not os.path.isfile(path):
    return []
  doc = _yaml(path) or []
  if not isinstance(doc, list):
    raise ValueError(f"{path} must be a list of facts.")
  return [dict(f) for f in doc if isinstance(f, dict)]

def cases_dir(explicit: str | None = None) -> str | None:
  return explicit or os.environ.get(CASES_DIR_ENV) or None

def find_files(case: dict, files_dir: str | None) -> tuple:
  found, problems = {}, []

  for name in case["pdfs"]:
    p = os.path.join(files_dir, name) if files_dir else None
    if not p or not os.path.isfile(p):
      problems.append(f"{name} is not in {files_dir or 'a folder named by TRACE_CASES_DIR'}.")
      continue

    want = case["sha256"].get(name)
    if want and CASE.sha256_of(p) != want:
      problems.append(f"{name} is not the file this case was made for (its hash differs).")
      continue
    found[name] = p
  return found, problems

def list_cases(root: str = HERE) -> list:
  if not os.path.isdir(root):
    return []
  return sorted(d for d in os.listdir(root) if os.path.isfile(os.path.join(root, d, "case.yaml")))

def _summary_header(ws):
  for r in range(1, 40):
    cells = [_text(ws.cell(r, c).value).lower() for c in range(1, 13)]
    d = next((i for i, x in enumerate(cells) if x.startswith("desc") or x in ("item", "particulars")), None)
    u = next((i for i, x in enumerate(cells) if x.startswith("unit")), None)
    q = next((i for i, x in enumerate(cells) if x.startswith("quant") or x.startswith("qty") or x.startswith("total")), None)
    if d is not None and u is not None and q is not None:
      return r, d + 1, u + 1, q + 1
  return None

def read_summary(data: bytes) -> tuple:
  import io
  import openpyxl
  wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
  lines, problems = [], []

  for ws in wb.worksheets:
    head = _summary_header(ws)
    if head is None:
      problems.append(f"Sheet {ws.title}: no header row with description, unit and quantity.")
      continue
    hr, cd, cu, cq = head
    section = ""

    for r in range(hr + 1, ws.max_row + 1):
      item = _text(ws.cell(r, cd).value)
      unit = _text(ws.cell(r, cu).value)
      qty = _number(ws.cell(r, cq).value)
      if not item:
        continue
      if not unit and qty is None:
        section = item
        continue
      if qty is None:
        continue

      lines.append(BillLine(row=r, section=f"{ws.title} / {section}" if len(wb.worksheets) > 1 else section, item=item, unit=unit, quantity=qty))
  return lines, problems

def read_workbook(path: str, fmt: str) -> tuple:
  with open(path, "rb") as fh:
    data = fh.read()
  if fmt == "WORKSHEET_MAP":
    try:
      return list(parse_bill(data).lines), []
    except ValueError as e:
      return [], [f"{os.path.basename(path)}: {e}"]
  return read_summary(data)

def _slug(text: str) -> str:
  return re.sub(r"[^A-Z0-9]+", "_", text.upper()).strip("_")[:40] or "ALL"

def draft_expected(lines, block_floor: dict | None = None) -> tuple:
  entries, left, used = [], [], {}

  for bl in lines:
    fam = bill_key(bl.item)
    if fam is None or fam == "LUMP_SUM":
      left.append({"item": bl.item, "unit": bl.unit, "quantity": bl.quantity, "section": bl.section})
      continue
    block = _slug(bl.section) if bl.section else "ALL"
    floor = (block_floor or {}).get(bl.section)
    base = f"{fam}/{floor or block}"
    n = used.get(base, 0) + 1
    used[base] = n

    key = base if n == 1 else f"{base}/{n}"
    e = {"key": key, "item": bl.item, "unit": bl.unit.lower(), "qs": round(float(bl.quantity), 3), "drawing": round(float(bl.quantity), 3), "class": "DRAWING",
      "tolerance_pct": T.default_for(fam), "note": bl.section.title() if bl.section else "", "family": fam}
    entries.append(e)
  return entries, left

FIELD_ORDER = ("key", "item", "unit", "qs", "drawing", "class", "tolerance_pct", "note", "family", "families", "floor", "grade", "match")

def dump_expected(entries: list, header: str = "") -> str:
  out = [("# " + h).rstrip() for h in header.splitlines()] if header else []
  for e in entries:
    first = True
    for f in FIELD_ORDER:
      if f not in e or e[f] in ("", [], None) and f not in ("drawing", "qs"):
        continue
      v = e[f]
      text = yaml.safe_dump({f: v}, default_flow_style=True, allow_unicode=True, width=1000).strip()
      text = text[1:-1].strip() if text.startswith("{") else text
      out.append(("- " if first else "  ") + text)
      first = False
  return "\n".join(out) + "\n"
