from __future__ import annotations
import argparse
import hashlib
import json
import os
import re
import sys
import openpyxl
import yaml

MM_FT, MM_IN = 304.8, 25.4
SFT_PER_M2 = 10.7639104
SOURCE_NOTE = "Drafted from the workbook; not used by any run until a person confirms it."
LINE_COLS = {"No": "D", "Length": "E", "Width": "F", "H": "G"}     

def _num(v) -> float | None:
  try:
    x = float(v)
  except (TypeError, ValueError):
    return None
  return x if x == x else None

def _seed(key: str):
  from app.engine.measure import rulebook as RB
  s = RB.SEED.get(key)
  return None if s is None else s.value

def _close(a, b, tol: float = 0.001) -> bool:
  if isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)):
    return len(a) == len(b) and all(_close(x, y, tol) for x, y in zip(a, b))
  if isinstance(a, dict) and isinstance(b, dict):
    return a.keys() == b.keys() and all(_close(a[k], b[k], tol) for k in a)
  fa, fb = _num(a), _num(b)
  if fa is None or fb is None:
    return a == b
  return abs(fa - fb) <= tol * max(abs(fa), abs(fb), 1e-9)

def _row(key, value, unit, cls, source, note, scope=None, seed_key=None, needs=None, seed_value=None) -> dict:
  sv = seed_value if seed_value is not None else _seed(seed_key or key)
  out = {"key": key, "scope": dict(scope or {}), "value": value, "unit": unit, "cls": cls, "origin": "QS_WORKBOOK", "status": "DRAFT", "confirmed_by": None,
    "source": source, "note": f"{note} {SOURCE_NOTE}".strip(), "needs": list(needs or [])}
  if sv is not None:
    out["seed"] = {"key": seed_key or key, "value": sv, "same": _close(value, sv)}
  return out

class _Draft:
  def __init__(self, path: str):
    self.path = path
    self.rules: list = []
    self.skipped: list = []
    self.wb = openpyxl.load_workbook(path)

  def src(self, sheet: str, cell: str, text: str, **more) -> dict:
    return {"workbook": os.path.basename(self.path), "sheet": sheet, "cell": cell, "text": text[:240], **more}

  def skip(self, what: str, where: str, why: str) -> None:
    self.skipped.append({"what": what, "where": where, "why": why})

  def constants(self) -> None:
    if "Constants and rules" not in self.wb.sheetnames:
      self.skip("Constants and rules", "-", "the workbook has no 'Constants and rules' tab")
      return
    ws = self.wb["Constants and rules"]
    used: set = set()
    
    for r in range(2, ws.max_row + 1):
      item, text, where, note = (str(ws.cell(r, c).value or "") for c in (1, 2, 3, 4))
      if not item:
        continue
      cell = f"B{r}"
      s = lambda extra=None: self.src("Constants and rules", cell, text, item=item, qs_cell=where, **(extra or {}))
      hit = False
      m1, m2 = re.search(r"(\d+(?:\.\d+)?)\s*cft per m3", text), re.search(r"(\d+(?:\.\d+)?)\s*sft per m2", text)
      if m1 or m2:
        hit = True
        if m1:
          self.rules.append(_row("cft_per_m3", float(m1.group(1)), "cft per m3", "STANDARD", s(), "The workbook types this conversion; the seed holds the exact one."))
        if m2:
          self.rules.append(_row("sft_per_m2", float(m2.group(1)), "sft per m2", "STANDARD", s(), "The workbook types this conversion; the seed holds the exact one."))
          
      m = re.search(r"\(#\)\^2\s*/\s*\(\s*(\d+(?:\.\d+)?)\s*x\s*(\d+(?:\.\d+)?)\s*\)\s*kg per ft", text)
      if m and "eighth" in note.lower():
        hit = True
        a, b = float(m.group(1)), float(m.group(2))
        per_ft = (MM_IN / 8.0) ** 2 * a * b * (MM_FT / 1000.0)
        weights = re.findall(r"#(\d)\s+(\d+\.\d+)", text)
        ok = all(abs(int(n) ** 2 / (a * b) - float(w)) <= 0.0006 for n, w in weights) if weights else None
        self.rules.append(_row("steel.unit_weight_divisor", round(per_ft, 2), "kg/m = d2 / value", "STANDARD", s(),
          f"kg per ft = d(eighths of an inch)^2 / ({a:g} x {b:g}); with d in mm that is d^2 / {per_ft:.2f} kg per m." + (" The workbook's own listed weights (#3 to #6) agree." if ok else ""), seed_key="steel.unit_weight_divisor"))
        
      m = re.search(r"(\d+(?:\.\d+)?)\s*kg per sft", text)
      if m:
        hit = True
        self.rules.append(_row("slab_steel_kg_per_m2", round(float(m.group(1)) * SFT_PER_M2, 4), "kg/m2", "JUDGEMENT", s(), f"{m.group(1)} kg per sft of slab area; {note}".strip()))
      m = re.search(r"(\d+(?:\.\d+)?)%\s*added", text)
      if m and "waste" in item.lower():
        hit = True
        self.rules.append(_row("steel_wastage_pct", float(m.group(1)), "%", "JUDGEMENT", s(), "Steel waste added in the Breakdown."))
      if "lintel length" in item.lower():
        hit = True
        pairs = [(float(a), float(b)) for a, b in re.findall(r"([\d.]+)(?:\s*ft)?(?:\s*opening)?\s*(?:->|→)\s*([\d.]+)(?:\s*ft)?", text)]
        stated = re.search(r"\+\s*([\d.]+)\s*ft bearing each side", text)
        eachs = sorted({round((b - a) / 2.0 * MM_FT, 1) for a, b in pairs})
        
        if pairs and len(eachs) == 1:
          why = ""
          if stated and abs(float(stated.group(1)) * MM_FT - eachs[0]) > 1.0:
            why = f" The wording says {stated.group(1)} ft each side; its {len(pairs)} worked examples add {eachs[0] / MM_FT * 2:g} ft in all, {eachs[0] / MM_FT:g} ft each side. The examples are used."
          self.rules.append(_row("lintel_bearing_mm", eachs[0], "mm", "JUDGEMENT", s({"examples": [list(p) for p in pairs]}), f"Bearing on the wall at each end of a lintel, from the opening and lintel lengths in the worked examples.{why}"))
        else:
          self.skip(item, f"Constants and rules!{cell}", "the worked examples of the lintel length do not give one bearing")
          
      m = re.search(r"([\d.]+)\s*ft\s*\(([\d.]+)\s*in\)", text)
      if "slab thickness" in item.lower() and m:
        hit = True
        self.rules.append(_row("slab_thickness_mm", round(float(m.group(2)) * MM_IN, 1), "mm", "JUDGEMENT", s(), f"Slab thickness the workbook types. {note}".strip()))
      m = re.search(r"(\d+)\s*x\s*(\d+)\s*x\s*(\d+(?:\.\d+)?)\s*ft deep", text)
      if "basement block" in item.lower() and m:
        hit = True
        self.rules.append(_row("basement_dig_mm", round(float(m.group(3)) * MM_FT, 1), "mm", "JUDGEMENT", s(), f"Depth of the basement block ({m.group(1)} x {m.group(2)} ft). {note}".strip()))
      m = re.search(r"(\d+(?:\.\d+)?)\s*ft deep", text)
      if "trench rule" in item.lower() and m:
        hit = True
        self.rules.append(_row("trench_depth_mm", round(float(m.group(1)) * MM_FT, 1), "mm", "JUDGEMENT", s(), "Depth of the trench the workbook types."))
        used.add("trench")
      if any(w in item.lower() for w in ("trench rule", "foundation strip", "brick foundation", "door frame")):
        hit = True          
      if not hit:
        self.skip(item, f"Constants and rules!{cell}", "no shape of the closed library or key of the rule book covers this rule")

  def line_items(self) -> list:
    if "Line items" not in self.wb.sheetnames:
      self.skip("Line items", "-", "the workbook has no 'Line items' tab")
      return []
    rows = list(self.wb["Line items"].iter_rows(values_only=True))
    return [(*r, i) for i, r in enumerate(rows[1:], start=2) if r and r[0]]     

  def cell(self, r, field: str) -> str:
    return f"{LINE_COLS[field]}{r[1]}"

  def strips(self, lines: list) -> None:
    groups: dict = {}
    for r in lines:
      if r[0] not in ("Excavation", "PCC 1;4;8", "Termite") or str(r[17] or "") != "line":
        continue
      w, h = _num(r[9]), _num(r[10])
      if w is None or not r[3]:
        continue
      groups.setdefault((r[0], str(r[3]), w, h), []).append(r)
      
    one_off = 0
    for (sheet, block, w, h), rs in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2])):
      if sheet == "Termite":
        continue
      if len(rs) < 2 or not h:
        one_off += len(rs)
        continue
      first = rs[0]
      
      self.rules.append(_row("formula.strip", {"formula": "strip", "params": {"width_mm": round(w * MM_FT, 1), "depth_mm": round(h * MM_FT, 1), "length_basis": "CENTRE"}}, None, "JUDGEMENT",
        self.src(sheet, self.cell(first, "Width"), f"{len(rs)} rows with width {w:g} ft and depth/thickness {h:g} ft", map_cell=f"Line items!J{first[-1]}", rows=len(rs), first_row=first[1]),
        f"{sheet}, {block}: {len(rs)} rows are the wall length x {w:g} ft x {h:g} ft. The length is taken along the centre line of the wall (the QS answered a centre line for brick; the workbook does not say for this item).",
        scope={"item": sheet, "block": block, "width_ft": w, "depth_ft": h}))
    terms = [k for k in groups if k[0] == "Termite"]
    
    if terms:
      n = sum(len(groups[k]) for k in terms)
      widths = sorted({k[2] for k in terms})
      self.skip("Termite strips", f"Termite!F ({n} rows, widths {', '.join(f'{x:g}' for x in widths)} ft)", "termite treatment is an area (width x length, in sft); the library's strip gives a volume (width x depth x length)")
    if one_off:
      self.skip("One-off strip widths", "Excavation, PCC", f"{one_off} rows have a width or depth that no other row shares (for example a width typed as =2.25+0.6); a rule needs a repeat")

  def frames(self, lines: list) -> None:
    rs = [r for r in lines if r[0] == "Door Frames" and r[13] and re.fullmatch(r"=\s*[\d.]+\s*\+\s*[\d.]+\s*\+\s*[\d.]+", str(r[13]))]
    if not rs:
      return
    seen: dict = {}
    heights = set()
    for r in rs:
      a, h1, h2 = (float(x) for x in re.findall(r"[\d.]+", str(r[13])))
      if h1 != h2:
        continue
      heights.add(h1)
      seen.setdefault((a, _num(r[9])), []).append(r)
    if len(heights) != 1 or not seen:
      self.skip("Door frame", "Door Frames", "the length formulas do not all add the same two heights to the width")
      return
  
    h = heights.pop()
    by_section: dict = {}
    for (a, sec), rr in seen.items():
      by_section.setdefault(sec, []).append((a, rr[0]))
    bands = []
    for sec, items in sorted(by_section.items(), key=lambda kv: max(a for a, _ in kv[1])):
      mx = max(a for a, _ in items)
      bands.append({"max_width_mm": round(mx * MM_FT, 1), "width_mm": round(sec * MM_FT, 1), "depth_mm": None})
    first = rs[0]
    
    self.rules.append(_row("formula.frame", {"formula": "frame", "params": {"add_mm": 0.0, "section_by_width": bands}}, None, "JUDGEMENT",
      self.src("Door Frames", self.cell(first, "Length"), f"{len(rs)} rows with length = width + {h:g} + {h:g} (ft)", map_cell=f"Line items!N{first[-1]}", rows=len(rs), first_row=first[1]),
      f"A door frame is (width + {h:g} ft + {h:g} ft) x a face width of {', '.join(f'{b['width_mm'] / MM_FT:g} ft up to a {b['max_width_mm'] / MM_FT:g} ft door' for b in bands)}. The two {h:g} ft are the door height, which the workbook types for every door ({h * MM_FT:g} mm). "
      "The workbook multiplies by the face width only and states no depth of the frame section; the library needs one, so the row cannot be confirmed until a person types it.",
      scope={"item": "Door Frames"}, needs=[f"section_by_width[{i}].depth_mm" for i in range(len(bands))]))

  def steps(self, lines: list) -> None:
    blocks: dict = {}
    for r in lines:
      if r[0] != "Brick Work" or not r[2] or "step" not in str(r[2]).lower() or str(r[17] or "") != "line":
        continue
      w, h = _num(r[9]), _num(r[10])
      if w is None or h is None:
        continue
      blocks.setdefault(str(r[2]), {}).setdefault((w, h), []).append(r)
    if not blocks:
      return
  
    tables = {"THICK": [], "THIN": []}
    cells = {"THICK": [], "THIN": []}
    for name in sorted(blocks, key=lambda n: (re.sub(r"\D", "", n) or "9", n)):
      groups = sorted((g for g in blocks[name].items() if len(g[1]) >= 2), key=lambda g: -g[0][0])
      
      if not groups:
        continue
      (w, h), rs = groups[0]
      tables["THICK"].append({"width_mm": round(w * MM_FT, 1), "height_mm": round(h * MM_FT, 1)})
      cells["THICK"].append(f"Brick Work!{self.cell(rs[0], 'Width')} ({name})")
      if len(groups) > 1:
        (w2, h2), rs2 = groups[1]
        tables["THIN"].append({"width_mm": round(w2 * MM_FT, 1), "height_mm": round(h2 * MM_FT, 1)})
        cells["THIN"].append(f"Brick Work!{self.cell(rs2[0], 'Width')} ({name})")
        
    for wall, text in (("THICK", "walls of 9 in and thicker (the group with the wider step in each block)"), ("THIN", "the 4.5 in walls (the group with the narrower step)")):
      if not tables[wall]:
        continue
      seed_h = (_seed("brick_step_heights_mm") or {}).get(wall)
      self.rules.append(_row("formula.steps", {"formula": "steps", "params": {"table": tables[wall]}}, None, "JUDGEMENT",
        self.src("Brick Work", cells[wall][0].split("!")[1].split(" ")[0], f"{len(tables[wall])} steps for {wall.lower()} walls", cells=cells[wall]),
        f"Steps of the brick foundation under {text}: " + ", ".join(f"{t['width_mm'] / MM_IN:g} in wide x {t['height_mm'] / MM_FT:g} ft high" for t in tables[wall]) + ". The step is a group of at least two rows with the same width and height; which group is the thin wall's is read from the narrower step, not from a label.",
        scope={"walls": wall}, seed_key="brick_step_heights_mm", seed_value=seed_h))
      row = self.rules[-1]
      if seed_h is not None:
        row["seed"] = {"key": "brick_step_heights_mm", "value": seed_h, "same": _close([t["height_mm"] for t in tables[wall]], seed_h), "compares": "the heights of the steps"}

def draft(path: str) -> dict:
  d = _Draft(path)
  d.constants()
  lines = d.line_items()
  d.strips(lines)
  d.frames(lines)
  d.steps(lines)
  with open(path, "rb") as fh:
    digest = hashlib.sha256(fh.read()).hexdigest()
  return {"origin": "QS_WORKBOOK", "status": "DRAFT", "workbook": os.path.basename(path), "sha256": digest, "rules": d.rules, "not_drafted": d.skipped,
    "note": "Every row is a DRAFT: confirmed_by is null and no run uses it until a person confirms it."}

def main(argv: list | None = None) -> int:
  ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
  ap.add_argument("workbook")
  ap.add_argument("--out", default=".")
  a = ap.parse_args(argv)
  doc = draft(a.workbook)
  os.makedirs(a.out, exist_ok=True)
  
  with open(os.path.join(a.out, "draft_rules.yaml"), "w", encoding="utf-8") as fh:
    fh.write("# DRAFT rules from " + doc["workbook"] + ". Nothing here is used by a run until a person confirms it.\n")
    yaml.safe_dump(doc, fh, sort_keys=False, allow_unicode=True, width=200)
  with open(os.path.join(a.out, "draft_rules.json"), "w", encoding="utf-8") as fh:
    json.dump(doc, fh, indent=1, ensure_ascii=False)
  for r in doc["rules"]:
    sd = r.get("seed")
    print(f"{r['key']:28} {json.dumps(r['value'])[:70]:70} {r['source']['sheet']}!{r['source']['cell']}" + ("" if sd is None else f"   seed {json.dumps(sd['value'])[:30]} {'same' if sd['same'] else 'DIFFERS'}") + (f"   needs {len(r['needs'])}" if r["needs"] else ""))
  print(f"{len(doc['rules'])} DRAFT rows, {len(doc['not_drafted'])} not drafted")
  
  for s in doc["not_drafted"]:
    print("  not drafted:", s["what"], "|", s["where"], "|", s["why"])
  return 0

if __name__ == "__main__":
  sys.exit(main())
