from __future__ import annotations
from app.modules.drawings_boq.steel_view.summary import label_of, BASIS_WORDS
from openpyxl.styles import PatternFill, Font, Alignment
import io
from openpyxl import Workbook

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
STATUS_WORDS = {"IN_RANGE": "Within the usual range", "BELOW_RANGE": "Below the usual range", "ABOVE_RANGE": "Above the usual range"}

def _sheet(ws, head, widths):
  ws.append(head)
  for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="top")
  for col, w in zip("ABCDEFGHIJKL", widths):
    ws.column_dimensions[col].width = w
  ws.freeze_panes = "A2"

def steel_xlsx(result: dict) -> bytes:
  wb = Workbook()
  ws = wb.active
  ws.title = "By diameter"
  _sheet(ws, ["Diameter (mm)", "Weight (kg)", "Length (m)", "Share %", "Marks", "Kinds of member"], (14, 14, 14, 10, 10, 60))
  
  for d in result.get("by_diameter", []):
    ws.append([d["dia_mm"], d["kg"], d["length_m"], d["share_pct"], d["members"], ", ".join(f"{label_of(f)} {k}" for f, k in d["by_family"].items())])
  wo = result.get("without_diameter") or {}
  if wo.get("members"):
    ws.append(["no diameter", wo["kg"], None, None, wo["members"], "Estimated from the size of the member; the bars are not known."])
  t = result.get("totals", {})
  ws.append([])
  ws.append(["Total (kg)", t.get("kg"), None, None, t.get("members"), f"Estimated {t.get('estimated_kg', 0)} kg ({t.get('estimated_share_pct', 0)} %); range {t.get('kg_low')} to {t.get('kg_high')} kg"])
  w = result.get("wastage") or {}
  ws.append(["Wastage", w.get("applied_pct") if w.get("applied_pct") is not None else "as run", None, None, None, f"The run added {w.get('as_run_pct')} % to bars read from the drawing."])

  f = wb.create_sheet("By kind of member")
  _sheet(f, ["Kind of member", "Marks", "Weight (kg)", "Low (kg)", "High (kg)", "Concrete (m3)", "kg per m3", "Usual range", "Result", "From members"], (22, 8, 14, 12, 12, 14, 12, 14, 24, 14))
  for r in result.get("by_family", []):
    f.append([r["label"], r["members"], r["kg"], r["kg_low"], r["kg_high"], r["concrete_m3"], r["kg_per_m3"], " - ".join(str(x) for x in r["range"]) if r.get("range") else "",
      STATUS_WORDS.get(r.get("status") or "", ""), r.get("intensity_from")])

  m = wb.create_sheet("Marks")
  _sheet(m, ["Kind of member", "Mark", "Floor", "Basis", "Count", "Weight (kg)", "Each (kg)", "Low (kg)", "High (kg)", "Concrete (m3)", "kg per m3", "Diameters (kg)"], (20, 18, 18, 36, 8, 12, 12, 12, 12, 14, 12, 40))
  for r in result.get("marks", []):
    m.append([label_of(r["family"]), r["tag"], r["level"], BASIS_WORDS.get(r["basis"], r["basis"]), r["count"], r["kg"], r["kg_each"], r["kg_low"], r["kg_high"], r["concrete_m3"], r["kg_per_m3"],
      ", ".join(f"{d} mm {k}" for d, k in (r.get("dia") or {}).items())])
    
  a = result.get("allowances")
  if a:
    al = wb.create_sheet("Allowances")
    _sheet(al, ["Item", "Weight (kg)", "Note"], (30, 14, 70))
    al.append(["Reinforcement bars", a.get("bars_kg"), "The weight of the bars, as on the other sheets."])
    al.append(["Binding wire", a.get("binding_wire_kg"), "A share of the weight of the bars; not drawn, so assumed."])
    al.append(["Chairs and spacers", a.get("chairs_kg"), "A rate per square metre of slab; not drawn, so assumed."])
    al.append(["Bars with allowances", a.get("with_bars_kg"), f"The allowances add {a.get('share_pct')} % to the bars."])
  for sheet in wb.worksheets:
    for row in sheet.iter_rows(min_row=2):

      for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue()
