from __future__ import annotations
from openpyxl.styles import Font, Alignment, PatternFill
import io
from openpyxl import Workbook

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
STATUS_WORDS = {"SAME": "No change", "CHANGED": "Changed", "ADDED": "New in the later run", "REMOVED": "Not in the later run"}
KIND_WORDS = {"ADDITION": "Addition", "OMISSION": "Omission", "NEW_ITEM": "New item", "ITEM_DROPPED": "Item dropped"}
AUDIENCE_SHEETS = (("ARCHITECT", "For the architect"), ("QS", "For the QS"), ("TRACE", "For Trace"))

def _sheet(ws, head, widths):
  ws.append(head)
  for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="top")
  for col, w in zip("ABCDEFGHIJKL", widths):
    ws.column_dimensions[col].width = w
  ws.freeze_panes = "A2"

def _save(wb) -> bytes:
  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue()

def _wrap(ws):
  for row in ws.iter_rows(min_row=2):
    for c in row:
      c.alignment = Alignment(wrap_text=True, vertical="top")

def difference_xlsx(result: dict) -> bytes:
  wb = Workbook()
  ws = wb.active
  ws.title = "Difference"
  _sheet(ws, ["Work item", "Name", "Unit", "Earlier run", "Later run", "Change", "Change %", "Result", "Source moved"], (16, 34, 8, 14, 14, 14, 10, 20, 14))
  
  for i in result.get("items", []):
    ws.append([i["code"], i.get("name") or "", i["unit"], i["a"], i["b"], i["delta"], i.get("pct"), STATUS_WORDS.get(i["status"], i["status"]), "yes" if i.get("source_moved") else ""])
  st = wb.create_sheet("By storey")
  _sheet(st, ["Work item", "Unit", "Storey", "Earlier run", "Later run", "Change"], (16, 8, 28, 14, 14, 14))
  for i in result.get("items", []):
    for s in i.get("storeys", []):
      st.append([i["code"], i["unit"], s["level"], s["a"], s["b"], s["delta"]])
  me = wb.create_sheet("By member")
  _sheet(me, ["Work item", "Unit", "Member", "Storey", "Result", "Earlier run", "Later run", "Change"], (16, 8, 24, 24, 12, 14, 14, 14))
  for i in result.get("items", []):
    for m in i.get("members", []):
      me.append([i["code"], i["unit"], m["label"], m.get("level") or "", m["status"].capitalize(), m["a"], m["b"], m["delta"]])
  why = wb.create_sheet("Why")
  _sheet(why, ["What differs", "Detail"], (16, 100))
  for c in result.get("causes", []):
    why.append([c["kind"].capitalize(), c["text"]])
  _wrap(why)
  return _save(wb)

def variation_xlsx(variation: dict) -> bytes:
  wb = Workbook()
  ws = wb.active
  ws.title = "Variation"
  labels = variation.get("labels") or {}
  ws.append([f"Variation statement: {labels.get('from', 'earlier run')} to {labels.get('to', 'later run')}"])
  ws["A1"].font = Font(bold=True, size=13)
  ws.append([variation.get("note", "")])
  ws.append([])
  head_row = ws.max_row + 1
  
  ws.append(["Work item", "Name", "Unit", "Kind", "Previous", "Revised", "Omission", "Addition", "Change %", "Basis"])
  for c in ws[head_row]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
  for col, w in zip("ABCDEFGHIJ", (16, 34, 8, 14, 14, 14, 14, 14, 10, 70)):
    ws.column_dimensions[col].width = w
  for l in variation.get("lines", []):
    ws.append([l["code"], l.get("name") or "", l["unit"], KIND_WORDS.get(l["kind"], l["kind"]), l["previous"], l["revised"], l["omission"] or None, l["addition"] or None, l.get("pct"), l.get("basis", "")])
  ws.freeze_panes = ws.cell(head_row + 1, 1)
  c = variation.get("counts") or {}
  sm = wb.create_sheet("Summary")
  
  for k, label in (("additions", "Items with an addition"), ("omissions", "Items with an omission"), ("new_items", "New items"), ("dropped_items", "Items dropped"),
      ("unchanged", "Items not changed"), ("below_threshold", f"Changes below {variation.get('min_change_pct')}% (not listed)")):
    sm.append([label, c.get(k)])
  sm.column_dimensions["A"].width = 44
  why = wb.create_sheet("Basis")
  _sheet(why, ["What differs", "Detail"], (16, 100))
  for x in variation.get("causes", []):
    why.append([x["kind"].capitalize(), x["text"]])
  _wrap(why)
  return _save(wb)

def fixlist_xlsx(fix: dict) -> bytes:
  wb = Workbook()
  first = True
  for audience, title in AUDIENCE_SHEETS:
    ws = wb.active if first else wb.create_sheet()
    first = False
    ws.title = title
    _sheet(ws, ["No", "Topic", "What is needed", "Where", "Why it matters", "How many", "Priority"], (6, 30, 70, 40, 70, 10, 10))
    
    n = 0
    for e in fix.get("items", []):
      if e["audience"] != audience:
        continue
      n += 1
      where = "; ".join(x for x in (", ".join(e.get("levels") or []), ("marks " + ", ".join(e["marks"])) if e.get("marks") else "") if x)
      ws.append([n, e["topic"], e["ask"], where, e["why"], e.get("count") or "", {"error": "High", "warning": "Medium", "info": "Low"}.get(e["severity"], e["severity"])])
    _wrap(ws)
  return _save(wb)
