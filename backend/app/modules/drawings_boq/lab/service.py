from __future__ import annotations
import io
from sqlalchemy.ext.asyncio import AsyncSession
from uuid import UUID
from sqlalchemy import select
from app.modules.drawings_boq.models import BuildingLevel, CalculationRun, QuantityLedger
from fastapi import HTTPException
from app.modules.drawings_boq.lab.bill import parse_bill
from app.modules.drawings_boq.lab.compare import CLASSES, DEFAULT_TOLERANCE_PCT, compare
from app.modules.drawings_boq.lab import steel_groups as SG
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

MAX_BILL_BYTES = 8 * 1024 * 1024

class BillCompareService:
  def __init__(self, session: AsyncSession):
    self.session = session

  async def _rows(self, org: UUID, run_id: UUID) -> list[dict]:
    run = (await self.session.execute(
      select(CalculationRun.id).where(CalculationRun.id == run_id, CalculationRun.organization_id == org))).scalar_one_or_none()
    if run is None:
      raise HTTPException(404, {"code": "RUN_NOT_FOUND", "message": "Calculation run not found."})
    result = await self.session.execute(
      select(QuantityLedger.formula_code, QuantityLedger.quantity_net, QuantityLedger.unit, QuantityLedger.trace)
      .where(QuantityLedger.run_id == run_id, QuantityLedger.organization_id == org))
    return [{"formula_code": f, "quantity": q, "unit": u, "trace": t} for f, q, u, t in result.all()]

  async def compare_bill(self, org: UUID, run_id: UUID, data: bytes, tolerance_pct: float | None, overrides: dict | None, classes: dict | None) -> dict:
    if len(data) > MAX_BILL_BYTES:
      raise HTTPException(413, {"code": "BILL_TOO_LARGE", "message": "The bill file is too large."})
    try:
      bill = parse_bill(data)
    except ValueError as e:
      raise HTTPException(422, {"code": str(e), "message": "The bill could not be read. It needs a summary sheet with Description, Unit and Quantity columns."})
    except Exception:
      raise HTTPException(422, {"code": "BILL_UNREADABLE", "message": "The bill file could not be opened as an Excel workbook."})
    tol = DEFAULT_TOLERANCE_PCT if tolerance_pct is None else max(0.0, min(float(tolerance_pct), 100.0))
    
    rows = await self._rows(org, run_id)
    out = compare(bill.lines, rows, tol, overrides, classes)
    out["bill_sheet"] = bill.sheet
    out["bill_sheets"] = list(bill.sheets)
    out["software_rows"] = len(rows)
    return out

  async def _steel_rows(self, org: UUID, run_id: UUID) -> list[dict]:
    """The steel rows of a run, with the kind of member, its mark and the floor of its level."""
    run = (await self.session.execute(
      select(CalculationRun.id).where(CalculationRun.id == run_id, CalculationRun.organization_id == org))).scalar_one_or_none()
    if run is None:
      raise HTTPException(404, {"code": "RUN_NOT_FOUND", "message": "Calculation run not found."})
  
    result = await self.session.execute(
      select(QuantityLedger.quantity_net, QuantityLedger.trace, QuantityLedger.warnings, BuildingLevel.name)
      .outerjoin(BuildingLevel, BuildingLevel.id == QuantityLedger.level_id)
      .where(QuantityLedger.run_id == run_id, QuantityLedger.organization_id == org, QuantityLedger.formula_code == "PDF_MEMBER_STEEL_WEIGHT"))
    out = []
    
    for kg, trace, warnings, level in result.all():
      src = (((trace or {}).get("inputs") or {}).get("steel_source") or {})
      out.append({"category": src.get("category"), "tag": src.get("tag"), "kg": float(kg or 0), "floor": SG.floor_of(level), "assumed": any(w == "STEEL_ASSUMED" or (isinstance(w, dict) and w.get("code") == "STEEL_ASSUMED") for w in (warnings or []))})
    return out

  async def compare_steel(self, org: UUID, run_id: UUID, data: bytes, tolerance_pct: float | None, classes: dict | None) -> dict:
    if len(data) > MAX_BILL_BYTES:
      raise HTTPException(413, {"code": "BILL_TOO_LARGE", "message": "The bill file is too large."})
    try:
      summary = SG.parse_steel_summary(data)
    except ValueError as e:
      raise HTTPException(422, {"code": str(e), "message": "The workbook has no Steel Summary sheet that can be read: it needs a heading row with Description and the floors (Basement, Ground Floor...)."})
    except Exception:
      raise HTTPException(422, {"code": "BILL_UNREADABLE", "message": "The bill file could not be opened as an Excel workbook."})
    tol = DEFAULT_TOLERANCE_PCT if tolerance_pct is None else max(0.0, min(float(tolerance_pct), 100.0))
    return SG.compare_steel(summary, await self._steel_rows(org, run_id), tol, classes)

def export_comparison(result: dict) -> bytes:
  wb = Workbook()
  ws = wb.active
  ws.title = "Comparison"
  head = ["Item", "Status", "Unit", "QS", "Software", "Difference", "Error %", "Class", "Note", "Bill rows"]
  ws.append(head)
  
  for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F3864")
  for x in result.get("lines", []):
    ws.append([x["label"], x["status"], x["unit"], x["qs"], x["software"], x["difference"], x["error_pct"], x.get("class") or x.get("suggested_class") or "", x.get("note", ""),
      "; ".join(f'{b["item"]} ({b["quantity"]})' for b in x.get("bill_rows", []))])
  for col, w in zip("ABCDEFGHIJ", (30, 14, 8, 14, 14, 14, 10, 18, 40, 60)):
    ws.column_dimensions[col].width = w
  m = wb.create_sheet("Measures")
  
  for k, v in (result.get("measures") or {}).items():
    m.append([k, v])
  m.append(["tolerance_pct", result.get("tolerance_pct")])
  if any(x.get("floors") for x in result.get("lines", [])):
    f = wb.create_sheet("By floor")
    f.append(["Item", "Floor", "QS", "Software"])
    
    for x in result["lines"]:
      for fl in x.get("floors") or []:
        f.append([x["label"], fl.get("label") or fl.get("floor"), fl.get("qs"), fl.get("software")])
  u = wb.create_sheet("Not matched")
  u.append(["Bill item", "Unit", "Quantity", "Section"])
  
  for x in result.get("unmapped", []):
    u.append([x["item"], x["unit"], x["quantity"], x["section"]])
  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue()

__all__ = ["BillCompareService", "export_comparison", "CLASSES"]
