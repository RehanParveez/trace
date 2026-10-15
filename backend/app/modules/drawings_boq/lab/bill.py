from __future__ import annotations
import io
import re
from dataclasses import dataclass
import openpyxl

SUMMARY_HINTS = ("breakdown", "summary")

@dataclass(frozen=True)
class BillLine:
  row: int
  section: str
  item: str
  unit: str
  quantity: float

@dataclass(frozen=True)
class Bill:
  sheet: str
  lines: tuple
  sheets: tuple

def _text(v) -> str:
  return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""

def _number(v):
  if isinstance(v, bool):
    return None
  if isinstance(v, (int, float)):
    return float(v)
  return None

def _summary_sheet(wb):
  for ws in wb.worksheets:
    top = " ".join(_text(c.value).lower() for row in ws.iter_rows(min_row=1, max_row=6, max_col=4) for c in row if c.value is not None)
    if any(h in top for h in SUMMARY_HINTS) or any(h in ws.title.lower() for h in SUMMARY_HINTS):
      return ws
  return wb.worksheets[0]

def _header_row(ws):
  for r in range(1, 15):
    cells = [_text(ws.cell(r, c).value).lower() for c in range(1, 7)]
    if any(x.startswith("desc") for x in cells) and any(x.startswith("unit") for x in cells) and any(x.startswith("quant") or x.startswith("total") for x in cells):
      return r, cells
  return None, []

def parse_bill(data: bytes) -> Bill:
  wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
  ws = _summary_sheet(wb)
  hr, head = _header_row(ws)
  if hr is None:
    raise ValueError("BILL_NO_HEADER")
  col_desc = next(i for i, x in enumerate(head) if x.startswith("desc")) + 1
  col_unit = next(i for i, x in enumerate(head) if x.startswith("unit")) + 1
  col_qty = next(i for i, x in enumerate(head) if x.startswith("quant") or x.startswith("total")) + 1
  section = ""
  lines = []
  
  for r in range(hr + 1, ws.max_row + 1):
    item = _text(ws.cell(r, col_desc).value)
    unit = _text(ws.cell(r, col_unit).value)
    qty = _number(ws.cell(r, col_qty).value)
    if not item:
      continue
    if not unit and qty is None:
      section = item
      continue
  
    if qty is None:
      continue
    lines.append(BillLine(row=r, section=section, item=item, unit=unit, quantity=qty))
  if not lines:
    raise ValueError("BILL_EMPTY")
  return Bill(sheet=ws.title, lines=tuple(lines), sheets=tuple(w.title for w in wb.worksheets))
