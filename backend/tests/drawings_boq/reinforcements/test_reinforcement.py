from __future__ import annotations
from decimal import Decimal as D
from app.engine.measure.rebar import _unique_mark
from app.modules.drawings_boq.rebar.rebar_parsing import rows_from_table
from openpyxl import load_workbook
from io import BytesIO
from app.modules.drawings_boq.export import build_bbs_xlsx

SNAP = {"version_no": 1, "purpose": "ISSUE", "content_hash": "a" * 64, "engine_version": "t",
  "rule_set_code": "T", "rule_set_version": 1, "convention_code": None}

def _mark(**kw):
  base = dict(level="GF", member="C1", trace={"member_mark": "C1"}, mark="M1", shape_code="STRAIGHT",
    shape_params={"A": "3000"}, dia_mm=D("16"), designation="16", grade="G60", count=10, spacing_mm=None,
    cut_len_mm=D("3000"), stock_len_mm=D("12000"), pieces=1, lap_count=0, lap_len_mm=None, total_len_m=D("30"),
    unit_weight_kg_m=D("1.579"), total_kg=D("47.370000"), provenance="SCHEDULE_IMPORT", confidence=D("0.9"),
    review_status="OK", warnings=[])
  base.update(kw)
  return base

def test_unique_mark_never_exceeds_column_width():
  used: dict = {}
  first, second = _unique_mark(used, "s", "X" * 60), _unique_mark(used, "s", "X" * 60)
  assert len(first) == 50 and len(second) <= 50 and first != second

def test_role_is_normalised_to_a_family_word():
  t = [["Member", "Bar Mark", "Role", "Dia", "Nos", "Length (mm)"],
      ["C1", "M1", "GF column", "16", "4", "3200"],
      ["Footing F2", "M2", "", "12", "6", "1500"]]
  assert [r["role"] for r in rows_from_table(t)] == ["COLUMN", "FOOTING"]

def test_guessed_length_unit_is_noted_and_lowers_confidence():
  t = [["Bar Mark", "Dia", "Nos", "Length"], ["M1", "16", "4", "3.2"], ["M2", "16", "4", "3200"]]
  a, b = rows_from_table(t)
  assert a["cut_len_mm"] == D("3200.000") and a["confidence"] == D("0.5") and a["notes"]
  assert b["confidence"] == D("0.9") and not b["notes"]

def test_bbs_workbook_totals_and_estimate_note():
  marks = [_mark(), _mark(mark="S1", shape_code="STIRRUP_RECT", shape_params={"A": "300", "B": "500"}, dia_mm=D("8"),
    designation="8", count=100, cut_len_mm=D("1664"), stock_len_mm=None, total_len_m=D("166.4"),
    unit_weight_kg_m=D("0.395"), total_kg=D("65.728000"))]
  
  wb = load_workbook(BytesIO(build_bbs_xlsx(SNAP, marks, {}, "Acme", D("120.5"))))
  sheet = [[c for c in r if c is not None] for r in wb["Bar Bending Schedule"].iter_rows(values_only=True)]
  assert any("120.500 kg" in str(r[0]) for r in sheet if r)
  assert sheet[-1] == ["TOTAL", 113.098]
  summary = [r for r in wb["Summary by diameter"].iter_rows(values_only=True)]
  assert summary[-1][0] == "TOTAL" and summary[-1][5] == 113.098