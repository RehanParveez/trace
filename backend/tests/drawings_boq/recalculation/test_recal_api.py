from __future__ import annotations
from datetime import datetime, timedelta, timezone
import pytest
from sqlalchemy import select, update
from app.modules.drawings_boq.models import CalculationRun, BOQVersion, ExportJob, QuantityLedger, RunDependencyEdge, RunElementState
from io import BytesIO
from uuid import uuid4
from app.core.config import settings
from app.modules.drawings_boq.boq_service import BOQEngineService
from openpyxl import load_workbook
from uuid import UUID
from app.modules.identity.enums import PermissionKey
from app.modules.drawings_boq.recalculation import calc_maintenance

P = "/api/v1/drawings-boq"

def _ok(r, code=200):
  assert r.status_code == code, r.text
  return r.json() if r.content else None

def _ago(**kw):
  return datetime.now(timezone.utc) - timedelta(**kw)

@pytest.fixture(autouse=True)
def _redis(monkeypatch, counter_redis):
  monkeypatch.setattr("app.modules.drawings_boq.recalculation.calc_limits.CalcLimits._client", lambda self: counter_redis)
  monkeypatch.setattr("app.modules.drawings_boq.recalculation.calc_cache.RunReadCache._client", lambda self: counter_redis)

async def _completed_run(db, org, project, make_run, lines=None):
  lines = lines or [("CONC-COL", "5", "m3", "G1"), ("CONC-SLAB", "9", "m3", "G2")]
  run = await make_run(org, project, lines)
  return run.id

async def test_run_metrics_endpoint(client, db_session, organization, actor_a, act_as, project_a, make_run):
  act_as(actor_a)
  rid = await _completed_run(db_session, organization, project_a, make_run)
  await db_session.execute(update(CalculationRun).where(CalculationRun.id == rid).values(
    started_at=_ago(seconds=60), completed_at=_ago(seconds=30),
    metrics={"timings_ms": {"total": 30000}, "counts": {"elements": 100}, "peak_rss_mb": 300}))
  db_session.expire_all()
  body = _ok(await client.get(f"{P}/calculation-runs/{rid}/metrics"))
  assert body["run_id"] == str(rid) and body["duration_ms"] == 30000 and body["mode"] == "FULL"
  assert {b["name"] for b in body["budgets"]} == {"calculation_time", "peak_memory"}
  assert all(b["ok"] for b in body["budgets"])

async def test_run_metrics_permissions_and_tenancy(client, db_session, organization, actor_a, actor_b, act_as, project_a,
  make_run, make_actor):
  rid = await _completed_run(db_session, organization, project_a, make_run)
  act_as(actor_b)
  assert (await client.get(f"{P}/calculation-runs/{rid}/metrics")).status_code == 404
  assert (await client.get(f"{P}/calculation-runs/{rid}/elements/{uuid4()}/impact")).status_code == 404
  nobody = await make_actor(organization, keys=[])
  act_as(nobody)
  assert (await client.get(f"{P}/calculation-runs/{rid}/metrics")).status_code == 403
  assert (await client.get(f"{P}/calculation-metrics")).status_code == 403
  assert (await client.get(f"{P}/calculation-usage")).status_code == 403

async def test_unknown_run_is_404(client, actor_a, act_as):
  act_as(actor_a)
  assert (await client.get(f"{P}/calculation-runs/{uuid4()}/metrics")).status_code == 404

async def test_organisation_metrics_and_usage_endpoints(client, db_session, organization, actor_a, act_as, project_a,
  make_run, monkeypatch):
  monkeypatch.setattr(settings, "calc_max_concurrent_runs_per_org", 3)
  monkeypatch.setattr(settings, "calc_run_rate_limit_per_minute", 5)
  act_as(actor_a)
  rid = await _completed_run(db_session, organization, project_a, make_run)
  db_session.expire_all()
  m = _ok(await client.get(f"{P}/calculation-metrics", params={"days": 7}))
  assert m["window_days"] == 7 and m["runs_total"] == 1 and m["runs_by_status"] == {"COMPLETED": 1}
  assert (await client.get(f"{P}/calculation-metrics", params={"days": 0})).status_code == 422
  u = _ok(await client.get(f"{P}/calculation-usage"))
  assert u["concurrent"] == {"active": 0, "limit": 3}
  assert {w["window_seconds"] for w in u["rate_limits"]} == {60, 3600}

async def test_organisation_metrics_do_not_leak_across_companies(client, db_session, organization, actor_a, actor_b, act_as,
  project_a, make_run):
  await _completed_run(db_session, organization, project_a, make_run)
  db_session.expire_all()
  act_as(actor_b)
  assert _ok(await client.get(f"{P}/calculation-metrics"))["runs_total"] == 0

async def test_run_request_accepts_force_full_and_verify_and_rejects_junk(client, actor_a, act_as, project_a):
  act_as(actor_a)
  r = await client.post(f"{P}/projects/{project_a.id}/calculation-runs", json={"force_full": "maybe"})
  assert r.status_code == 422
  r = await client.post(f"{P}/projects/{uuid4()}/calculation-runs", json={"force_full": True, "verify": True})
  assert r.status_code == 404

async def test_impact_endpoint(client, db_session, organization, actor_a, act_as, project_a, make_run):
  act_as(actor_a)
  rid = await _completed_run(db_session, organization, project_a, make_run)
  ledger = (await db_session.execute(select(QuantityLedger).where(QuantityLedger.run_id == rid)
    .order_by(QuantityLedger.work_item_code))).scalars().all()
  oid = organization.id
  els = [r.element_id for r in ledger]
  for i, r in enumerate(ledger):
    db_session.add(RunElementState(id=uuid4(), organization_id=oid, run_id=rid, element_key=f"g:K{i}",
      element_id=r.element_id, solid_id=r.solid_id, alloc_hash="h" * 64, participating=True, warnings=[], owned_mm3=1.0))
  await db_session.flush()
  db_session.add(RunDependencyEdge(id=uuid4(), organization_id=oid, run_id=rid, key_a="g:K0", key_b="g:K1",
    element_a_id=els[0], element_b_id=els[1], overlap_mm3=1000.0))
  await db_session.execute(update(CalculationRun).where(CalculationRun.id == rid).values(state_available=True))
  db_session.expire_all()
  body = _ok(await client.get(f"{P}/calculation-runs/{rid}/elements/{els[0]}/impact"))
  assert body["touching"][0]["element_id"] == str(els[1]) and body["touching"][0]["overlap_mm3"] == 1000.0
  assert len(body["affected_ledger"]) == 2 and body["truncated"] is False
  r = await client.get(f"{P}/calculation-runs/{rid}/elements/{uuid4()}/impact")
  assert r.status_code == 404 and "ELEMENT_NOT_IN_RUN" in r.text

async def test_finished_run_pages_are_cached_and_scoped_to_the_company(client, db_session, organization, actor_a, actor_b,
  act_as, project_a, make_run, counter_redis):
  act_as(actor_a)
  rid = await _completed_run(db_session, organization, project_a, make_run)
  first = _ok(await client.get(f"{P}/calculation-runs/{rid}/ledger", params={"limit": 10}))
  assert len(first) == 2
  assert any(k.startswith("trace:calc:read:") and ":ledger:" in k for k in counter_redis.store)
  await db_session.execute(update(QuantityLedger).where(QuantityLedger.run_id == rid).values(work_item_code="CHANGED"))
  second = _ok(await client.get(f"{P}/calculation-runs/{rid}/ledger", params={"limit": 10}))
  assert second == first                                   
  other = _ok(await client.get(f"{P}/calculation-runs/{rid}/ledger", params={"limit": 5}))
  assert {r["work_item_code"] for r in other} == {"CHANGED"}   
  act_as(actor_b)
  assert (await client.get(f"{P}/calculation-runs/{rid}/ledger", params={"limit": 10})).status_code == 404

async def test_running_run_pages_are_never_cached(client, db_session, organization, actor_a, act_as, project_a, make_run,
  counter_redis):
  act_as(actor_a)
  rid = await _completed_run(db_session, organization, project_a, make_run)
  await db_session.execute(update(CalculationRun).where(CalculationRun.id == rid).values(status="RUNNING"))
  db_session.expire_all()
  _ok(await client.get(f"{P}/calculation-runs/{rid}/solids", params={"limit": 10}))
  assert not [k for k in counter_redis.store if k.startswith("trace:calc:read:")]

async def test_cursor_paging_still_works_through_the_cache(client, db_session, organization, actor_a, act_as, project_a,
  make_run):
  act_as(actor_a)
  rid = await _completed_run(db_session, organization, project_a, make_run,
    [("CONC-COL", "1", "m3", "A"), ("CONC-SLAB", "2", "m3", "B"), ("BRICK", "3", "m3", "C")])
  seen, after = [], None
  for _ in range(5):
    params = {"limit": 1, **({"after": after} if after else {})}
    r = await client.get(f"{P}/calculation-runs/{rid}/ledger", params=params)
    assert r.status_code == 200, r.text
    seen += [row["id"] for row in r.json()]
    after = r.headers.get("X-Next-Cursor")
    if not after:
      break
  assert len(seen) == 3 and len(set(seen)) == 3

@pytest.fixture
def fake_storage(monkeypatch):
  files = {}
  def upload(key, fileobj, content_type=None):
    files[key] = fileobj.read()
  monkeypatch.setattr("app.shared.storage.upload_fileobj", upload)
  monkeypatch.setattr("app.shared.storage.download_bytes", lambda key: files[key])
  monkeypatch.setattr("app.shared.storage.delete_object", lambda key: files.pop(key))
  return files

@pytest.fixture
def queued(monkeypatch):
  sent = []
  monkeypatch.setattr("app.modules.drawings_boq.recalculation.export_task.render_export_task.delay", lambda job_id: sent.append(job_id))
  return sent

async def _approved(db, org, actor, project, make_run, make_book):
  await make_book(None, f"CSR{uuid4().hex[:4]}", rates=[("CONC-COL", "m3", 20000), ("CONC-SLAB", "m3", 18000)])
  run = await make_run(org, project, [("CONC-COL", "10", "m3", "G1"), ("CONC-SLAB", "4", "m3", "G2")])
  rid = run.id
  engine = BOQEngineService(db)
  vid = (await engine.build_from_run(org.id, rid, actor.user_id))["boq_version_id"]
  await engine.submit_for_review(org.id, vid, actor.user_id)
  await db.refresh(await db.get(BOQVersion, vid))
  await engine.approve_version(org.id, vid, actor.user_id)
  return vid

async def test_async_export_end_to_end(client, db_session, organization, actor_a, act_as, project_a, make_run, make_book,
  fake_storage, queued):
  act_as(actor_a)
  vid = await _approved(db_session, organization, actor_a, project_a, make_run, make_book)
  r = await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"fmt": "xlsx", "mode": "async"})
  job = _ok(r, 202)
  assert job["status"] == "QUEUED" and job["kind"] == "CONTRACT_BOQ" and job["format"] == "XLSX"
  assert r.headers["location"].endswith(f"/export-jobs/{job['id']}")
  assert queued == [job["id"]]
  assert (await client.get(f"{P}/export-jobs/{job['id']}/download")).status_code == 409
  
  assert await BOQEngineService(db_session).run_export_job(UUID(job["id"])) == "done"
  db_session.expire_all()
  done = _ok(await client.get(f"{P}/export-jobs/{job['id']}"))
  assert done["status"] == "SUCCEEDED" and done["file_size_bytes"] > 0 and done["error_code"] is None
  dl = await client.get(f"{P}/export-jobs/{job['id']}/download")
  assert dl.status_code == 200 and "attachment" in dl.headers["content-disposition"]
  assert load_workbook(BytesIO(dl.content)).sheetnames
  again = await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"fmt": "xlsx", "mode": "async"})
  assert again.status_code == 200 and again.json()["id"] == job["id"]
  assert queued == [job["id"]]
  assert await BOQEngineService(db_session).run_export_job(UUID(job["id"])) == "skipped"

async def test_sync_mode_and_threshold_routing(client, db_session, organization, actor_a, act_as, project_a, make_run, make_book,
  fake_storage, queued, monkeypatch):
  act_as(actor_a)
  vid = await _approved(db_session, organization, actor_a, project_a, make_run, make_book)
  r = await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"fmt": "xlsx"})
  assert r.status_code == 200 and load_workbook(BytesIO(r.content)).sheetnames        
  monkeypatch.setattr(settings, "export_async_item_threshold", 1)
  r = await client.get(f"{P}/boq-versions/{vid}/exports/PROCUREMENT", params={"fmt": "xlsx"})
  assert r.status_code == 202 and r.json()["kind"] == "PROCUREMENT"                  
  r = await client.get(f"{P}/boq-versions/{vid}/exports/BBS", params={"fmt": "xlsx", "mode": "sync"})
  assert r.status_code != 202                                                           
  assert (await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"mode": "later"})).status_code == 422

async def test_async_export_refusals_come_before_queuing(client, db_session, organization, actor_a, act_as, project_a,
  make_run, queued):
  act_as(actor_a)
  run = await make_run(organization, project_a, [("CONC-COL", "10", "m3", "G1")])
  vid = (await BOQEngineService(db_session).build_from_run(organization.id, run.id, actor_a.user_id))["boq_version_id"]
  r = await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"fmt": "xlsx", "mode": "async"})
  assert r.status_code == 409 and "EXPORT_REQUIRES_APPROVED" in r.text
  assert (await client.get(f"{P}/boq-versions/{vid}/exports/PROCUREMENT", params={"fmt": "pdf", "mode": "async"})).status_code == 422
  assert queued == []

async def test_failed_render_marks_the_job_failed(client, db_session, organization, actor_a, act_as, project_a, make_run, make_book,
  fake_storage, queued, monkeypatch):
  act_as(actor_a)
  vid = await _approved(db_session, organization, actor_a, project_a, make_run, make_book)
  job = _ok(await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"fmt": "xlsx", "mode": "async"}), 202)
  async def boom(*a, **k):
    raise RuntimeError("renderer crashed")
  monkeypatch.setattr(BOQEngineService, "_render_export", boom)
  
  assert await BOQEngineService(db_session).run_export_job(UUID(job["id"])) == "failed"
  db_session.expire_all()
  body = _ok(await client.get(f"{P}/export-jobs/{job['id']}"))
  assert body["status"] == "FAILED" and body["error_code"] == "EXPORT_FAILED" and "renderer crashed" in body["error_message"]
  assert (await client.get(f"{P}/export-jobs/{job['id']}/download")).status_code == 409

  monkeypatch.undo()
  again = await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"fmt": "xlsx", "mode": "async"})
  assert again.status_code == 202 and again.json()["id"] != job["id"]

async def test_export_jobs_are_private_to_their_company_and_permission(client, db_session, organization, actor_a, actor_b,
  act_as, project_a, make_run, make_book, make_actor, fake_storage, queued):
  act_as(actor_a)
  vid = await _approved(db_session, organization, actor_a, project_a, make_run, make_book)
  job = _ok(await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"fmt": "xlsx", "mode": "async"}), 202)
  act_as(actor_b)
  assert (await client.get(f"{P}/export-jobs/{job['id']}")).status_code == 404
  assert (await client.get(f"{P}/export-jobs/{job['id']}/download")).status_code == 404
  assert (await client.get(f"{P}/boq-versions/{vid}/exports/CONTRACT_BOQ", params={"mode": "async", "fmt": "xlsx"})).status_code == 404
  viewer = await make_actor(organization, keys=[PermissionKey.DRAWING_READ])
  act_as(viewer)
  assert (await client.get(f"{P}/export-jobs/{job['id']}")).status_code == 403
  assert (await client.get(f"{P}/export-jobs/{uuid4()}")).status_code in (403, 404)

async def test_lost_and_dead_export_jobs_are_recovered_and_files_expire(db_session, organization, actor_a, project_a,
  make_run, make_book, fake_storage, queued):
  vid = await _approved(db_session, organization, actor_a, project_a, make_run, make_book)
  svc = BOQEngineService(db_session)
  how, job = await svc.request_export(organization.id, vid, "CONTRACT_BOQ", "xlsx", actor_a.user_id, mode="async")
  jid = job.id
  await db_session.execute(update(ExportJob).where(ExportJob.id == jid).values(created_at=_ago(hours=1)))
  db_session.expire_all()
  sent = []
  assert (await calc_maintenance.recover_exports(db_session, send=sent.append))["resent"] == 1 and sent == [jid]

  await db_session.execute(update(ExportJob).where(ExportJob.id == jid).values(status="RUNNING", started_at=_ago(hours=3)))
  db_session.expire_all()
  assert (await calc_maintenance.recover_exports(db_session, send=sent.append))["failed"] == 1
  failed = await db_session.get(ExportJob, jid)
  assert failed.status == "FAILED" and failed.error_code == "EXPORT_STALE"
  
  fake_storage["k1"] = b"x"
  await db_session.execute(update(ExportJob).where(ExportJob.id == jid).values(
    status="SUCCEEDED", storage_key="k1", finished_at=_ago(days=int(settings.export_retention_days) + 1)))
  db_session.expire_all()
  removed = await calc_maintenance.expire_export_files(db_session, delete_object=fake_storage.pop)
  assert removed == {"removed": 1} and "k1" not in fake_storage
  db_session.expire_all()
  assert (await db_session.get(ExportJob, jid)).storage_key is None