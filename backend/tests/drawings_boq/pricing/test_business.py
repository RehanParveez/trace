from __future__ import annotations
from io import BytesIO
from uuid import uuid4
import pytest
from openpyxl import load_workbook
from app.modules.drawings_boq.boq_service import BOQEngineService
from app.modules.drawings_boq.models import BOQVersion
from app.modules.identity.enums import PermissionKey

P = "/api/v1/drawings-boq"

def _ok(r, code=200):
  assert r.status_code == code, r.text
  return r.json() if r.content else None

def _err(r, status, code):
  assert r.status_code == status, r.text
  assert r.json()["code"] == code if "code" in r.json() else code in r.text

async def _book(client, code="API-BOOK", **extra):
  return _ok(await client.post(f"{P}/rate-books", json={"code": code, "name": "API book", **extra}), 201)

async def _item(client, book_id, code="CONC-COL", unit="m3", rate="20000"):
  return _ok(await client.post(f"{P}/rate-books/{book_id}/items", json={"work_item_code": code, "unit": unit, "rate": rate}), 201)

async def test_rate_book_crud_and_lifecycle(client, actor_a, act_as):
  act_as(actor_a)
  b = await _book(client, edition="2026", description="x")
  assert b["status"] == "DRAFT" and b["code"] == "API-BOOK" and b["immutable_version"] == 1 and b["item_count"] == 0
  assert _ok(await client.get(f"{P}/rate-books/{b['id']}"))["id"] == b["id"]
  upd = _ok(await client.patch(f"{P}/rate-books/{b['id']}", json={"name": "Renamed"}))
  assert upd["name"] == "Renamed"
  _err(await client.post(f"{P}/rate-books/{b['id']}/publish"), 422, "RATE_BOOK_EMPTY")
  await _item(client, b["id"])
  pub = _ok(await client.post(f"{P}/rate-books/{b['id']}/publish"))
  assert pub["status"] == "ACTIVE" and pub["content_hash"]
  _err(await client.patch(f"{P}/rate-books/{b['id']}", json={"name": "x"}), 409, "RATE_BOOK_IMMUTABLE")
  _err(await client.delete(f"{P}/rate-books/{b['id']}"), 409, "RATE_BOOK_IMMUTABLE")
  nv = _ok(await client.post(f"{P}/rate-books/{b['id']}/new-version"), 201)
  assert nv["status"] == "DRAFT" and nv["immutable_version"] == 2 and nv["item_count"] == 1
  _ok(await client.delete(f"{P}/rate-books/{nv['id']}"), 204)
  _err(await client.get(f"{P}/rate-books/{nv['id']}"), 404, "RATE_BOOK_NOT_FOUND")
  assert _ok(await client.post(f"{P}/rate-books/{b['id']}/archive"))["status"] == "ARCHIVED"

async def test_rate_book_listing_filters_and_cursor(client, actor_a, act_as, make_book):
  act_as(actor_a)
  await make_book(None, "SYS-CSR", rates=[("A", "m3", 1)])
  for n in range(3):
    await _book(client, f"MINE-{n}")
  r = await client.get(f"{P}/rate-books", params={"owner": "org", "limit": 2})
  page = _ok(r)
  assert [x["code"] for x in page] == ["MINE-0", "MINE-1"] and r.headers["X-Next-Cursor"] == page[-1]["id"]
  rest = _ok(await client.get(f"{P}/rate-books", params={"owner": "org", "limit": 2, "after": r.headers["X-Next-Cursor"]}))
  assert [x["code"] for x in rest] == ["MINE-2"]
  sysrows = _ok(await client.get(f"{P}/rate-books", params={"owner": "system", "status": "ACTIVE"}))
  assert [x["code"] for x in sysrows] == ["SYS-CSR"] and sysrows[0]["organization_id"] is None
  assert _ok(await client.get(f"{P}/rate-books", params={"status": "ARCHIVED"})) == []
  assert (await client.get(f"{P}/rate-books", params={"limit": 0})).status_code == 422
  assert (await client.get(f"{P}/rate-books", params={"status": "BOGUS"})).status_code == 422

async def test_create_book_validation_and_copy_from(client, actor_a, act_as, make_book):
  act_as(actor_a)
  assert (await client.post(f"{P}/rate-books", json={"name": "no code"})).status_code == 422
  assert (await client.post(f"{P}/rate-books", json={"code": "X", "name": "x", "currency": "RUPEE"})).status_code == 422
  system = await make_book(None, "CSR", rates=[("CONC-COL", "m3", 21000)])
  copy = _ok(await client.post(f"{P}/rate-books", json={"code": "COPY", "name": "Copy", "copy_from_id": str(system.id)}), 201)
  assert copy["item_count"] == 1
  _err(await client.post(f"{P}/rate-books", json={"code": "C2", "name": "x", "copy_from_id": str(uuid4())}), 404, "RATE_BOOK_NOT_FOUND")

async def test_rate_item_endpoints(client, actor_a, act_as):
  act_as(actor_a)
  b = await _book(client)
  it = await _item(client, b["id"], "conc-col", "cum")
  assert (it["work_item_code"], it["unit"], it["rate"]) == ("CONC-COL", "m3", "20000.00")
  _err(await client.post(f"{P}/rate-books/{b['id']}/items", json={"work_item_code": "CONC-COL", "unit": "m3", "rate": "1"}), 409, "RATE_ITEM_EXISTS")
  assert (await client.post(f"{P}/rate-books/{b['id']}/items", json={"work_item_code": "A", "unit": "m3", "rate": "-1"})).status_code == 422
  bulk = _ok(await client.post(f"{P}/rate-books/{b['id']}/items/bulk", json={"items": [
    {"work_item_code": "CONC-COL", "unit": "m3", "rate": "21000"}, {"work_item_code": "STEEL", "unit": "kg", "rate": "280", "trade": "Steel"}]}))
  assert bulk == {"created": 1, "updated": 1, "total": 2}
  listed = _ok(await client.get(f"{P}/rate-books/{b['id']}/items", params={"q": "steel"}))
  assert [x["work_item_code"] for x in listed] == ["STEEL"]
  assert [x["work_item_code"] for x in _ok(await client.get(f"{P}/rate-books/{b['id']}/items", params={"trade": "steel"}))] == ["STEEL"]
  page = await client.get(f"{P}/rate-books/{b['id']}/items", params={"limit": 1})
  assert len(_ok(page)) == 1 and page.headers["X-Next-Cursor"]
  upd = _ok(await client.patch(f"{P}/rate-items/{listed[0]['id']}", json={"rate": "290", "is_active": False}))
  assert upd["rate"] == "290.00" and upd["is_active"] is False
  assert [x["work_item_code"] for x in _ok(await client.get(f"{P}/rate-books/{b['id']}/items"))] == ["CONC-COL"]
  assert len(_ok(await client.get(f"{P}/rate-books/{b['id']}/items", params={"include_inactive": True}))) == 2
  _ok(await client.delete(f"{P}/rate-items/{listed[0]['id']}"), 204)
  _err(await client.patch(f"{P}/rate-items/{uuid4()}", json={"rate": "1"}), 404, "RATE_ITEM_NOT_FOUND")
  _err(await client.delete(f"{P}/rate-items/{uuid4()}"), 404, "RATE_ITEM_NOT_FOUND")

async def test_csv_import_endpoint(client, actor_a, act_as):
  act_as(actor_a)
  b = await _book(client)
  csv = b"work_item_code,unit,rate,trade\nCONC-COL,cum,\"21,000\",Concrete\nSTEEL,kg,280,Steel\n"
  out = _ok(await client.post(f"{P}/rate-books/{b['id']}/import", files={"file": ("rates.csv", csv, "text/csv")}))
  assert out == {"created": 2, "updated": 0, "total": 2}
  bad = await client.post(f"{P}/rate-books/{b['id']}/import", files={"file": ("rates.csv", b"work_item_code,unit,rate\nA,m3,zzz\n", "text/csv")})
  _err(bad, 422, "RATE_IMPORT_INVALID")
  big = await client.post(f"{P}/rate-books/{b['id']}/import", files={"file": ("big.csv", b"x" * (5 * 1024 * 1024 + 10), "text/csv")})
  _err(big, 413, "RATE_IMPORT_TOO_LARGE")
  assert (await client.post(f"{P}/rate-books/{b['id']}/import")).status_code == 422

async def test_escalation_endpoints(client, actor_a, act_as):
  act_as(actor_a)
  b = await _book(client)
  e = _ok(await client.post(f"{P}/rate-books/{b['id']}/escalations", json={"trade_scope": "all", "effective_from": "2026-07-01", "factor": "1.05", "note": "Q3"}), 201)
  assert e["trade_scope"] == "ALL" and float(e["factor"]) == 1.05
  _err(await client.post(f"{P}/rate-books/{b['id']}/escalations", json={"effective_from": "2026-07-01", "factor": "1.1"}), 409, "ESCALATION_EXISTS")
  assert (await client.post(f"{P}/rate-books/{b['id']}/escalations", json={"effective_from": "2026-08-01", "factor": "0"})).status_code == 422
  assert [x["id"] for x in _ok(await client.get(f"{P}/rate-books/{b['id']}/escalations"))] == [e["id"]]
  _ok(await client.delete(f"{P}/rate-escalations/{e['id']}"), 204)
  _err(await client.delete(f"{P}/rate-escalations/{e['id']}"), 404, "ESCALATION_NOT_FOUND")

async def test_analysis_endpoints(client, actor_a, act_as):
  act_as(actor_a)
  b = await _book(client)
  body = {"code": "an-1", "work_item_code": "conc-col", "description": "Concrete", "unit": "m3", "overhead_pct": "10", "profit_pct": "15",
    "components": [{"component_type": "MATERIAL", "description": "Cement", "unit": "bag", "coefficient": "8", "rate_source": "DIRECT", "unit_rate": "1450"},
      {"component_type": "LABOUR", "description": "Mason", "unit": "day", "coefficient": "1.5", "rate_source": "DIRECT", "unit_rate": "2200"}]}
  a = _ok(await client.post(f"{P}/rate-books/{b['id']}/analyses", json=body), 201)
  assert a["code"] == "AN-1" and len(a["components"]) == 2 and a["computed_rate"] is None
  _err(await client.post(f"{P}/rate-books/{b['id']}/analyses", json=body), 409, "RATE_ANALYSIS_EXISTS")
  bad = dict(body, code="bad", components=[{"component_type": "MATERIAL", "description": "x", "unit": "u", "coefficient": "1", "rate_source": "DIRECT"}])
  assert (await client.post(f"{P}/rate-books/{b['id']}/analyses", json=bad)).status_code == 422
  assert [x["id"] for x in _ok(await client.get(f"{P}/rate-books/{b['id']}/analyses"))] == [a["id"]]
  assert _ok(await client.get(f"{P}/rate-analyses/{a['id']}"))["id"] == a["id"]
  bd = _ok(await client.get(f"{P}/rate-analyses/{a['id']}/breakdown"))
  assert bd["cost"] == "14900.00" and bd["overhead"] == "1490.00" and bd["profit"] == "2458.50" and bd["rate"] == "18848.50"
  comp = _ok(await client.post(f"{P}/rate-analyses/{a['id']}/compute"))
  assert comp["computed_rate"] == "18848.50"
  applied = _ok(await client.post(f"{P}/rate-analyses/{a['id']}/apply"))
  assert applied["work_item_code"] == "CONC-COL" and applied["rate"] == "18848.50" and applied["analysis_id"] == a["id"]
  up = _ok(await client.patch(f"{P}/rate-analyses/{a['id']}", json={"profit_pct": "0"}))
  assert up["computed_rate"] is None and up["profit_pct"] in ("0", "0.000", "0.0")
  _ok(await client.delete(f"{P}/rate-analyses/{a['id']}"), 204)
  _err(await client.get(f"{P}/rate-analyses/{a['id']}"), 404, "RATE_ANALYSIS_NOT_FOUND")
  _err(await client.post(f"{P}/rate-analyses/{uuid4()}/compute"), 404, "RATE_ANALYSIS_NOT_FOUND")

async def test_override_endpoints(client, actor_a, act_as, project_a):
  act_as(actor_a)
  url = f"{P}/projects/{project_a.id}/rate-overrides"
  o = _ok(await client.post(url, json={"work_item_code": "conc-col", "unit": "cum", "rate": "19000", "reason": "Quote"}), 201)
  assert (o["work_item_code"], o["unit"], o["revoked_at"]) == ("CONC-COL", "m3", None)
  _err(await client.post(url, json={"work_item_code": "CONC-COL", "unit": "m3", "rate": "1", "reason": "dup"}), 409, "RATE_OVERRIDE_EXISTS")
  assert (await client.post(url, json={"work_item_code": "A", "unit": "m3", "rate": "1"})).status_code == 422
  assert [x["id"] for x in _ok(await client.get(url))] == [o["id"]]
  assert (await client.post(f"{P}/rate-overrides/{o['id']}/revoke", json={})).status_code == 422
  rv = _ok(await client.post(f"{P}/rate-overrides/{o['id']}/revoke", json={"reason": "Price changed"}))
  assert rv["revoked_at"] and rv["revoke_reason"] == "Price changed"
  _err(await client.post(f"{P}/rate-overrides/{o['id']}/revoke", json={"reason": "again"}), 409, "RATE_OVERRIDE_REVOKED")
  assert _ok(await client.get(url)) == []
  assert len(_ok(await client.get(url, params={"include_revoked": True}))) == 1
  _err(await client.get(f"{P}/projects/{uuid4()}/rate-overrides"), 404, "PROJECT_NOT_FOUND")
  _err(await client.post(f"{P}/rate-overrides/{uuid4()}/revoke", json={"reason": "x"}), 404, "RATE_OVERRIDE_NOT_FOUND")

async def _built(db_session, org, actor, project, make_run, lines):
  run = await make_run(org, project, lines)
  return (await BOQEngineService(db_session).build_from_run(org.id, run.id, actor.user_id))["boq_version_id"]

async def test_price_summary_explain_and_diff_endpoints(client, db_session, organization, actor_a, act_as, project_a, make_run, make_book):
  act_as(actor_a)
  await make_book(None, "CSR", rates=[("CONC-COL", "m3", 20000), ("STEEL-60", "kg", 280)])
  v1 = await _built(db_session, organization, actor_a, project_a, make_run, [("CONC-COL", "10", "m3", "G1"), ("STEEL-60", "100", "kg", "G2")])
  summary = _ok(await client.get(f"{P}/boq-versions/{v1}/pricing"))
  assert summary["item_count"] == 2 and summary["unpriced_count"] == 0 and summary["total"] == "228000.00"
  assert summary["by_source"] == {"RATE_BOOK": 2}
  out = _ok(await client.post(f"{P}/boq-versions/{v1}/price"))
  assert out["priced"] == 2 and out["unpriced"] == 0 and out["total"] == "228000.00" and out["rate_books"][0]["code"] == "CSR"
  out2 = _ok(await client.post(f"{P}/boq-versions/{v1}/price", json={"as_of": "2030-01-01", "overwrite_manual": True}))
  assert out2["as_of"] == "2030-01-01"
  _err(await client.post(f"{P}/boq-versions/{v1}/price", json={"rate_book_ids": [str(uuid4())]}), 404, "RATE_BOOK_NOT_FOUND")
  _err(await client.post(f"{P}/boq-versions/{uuid4()}/price"), 404, "BOQ_VERSION_NOT_FOUND")
  _err(await client.get(f"{P}/boq-versions/{uuid4()}/pricing"), 404, "BOQ_VERSION_NOT_FOUND")

  item_id = str((await db_session.execute(__import__("sqlalchemy").text(
    "select id from boq_items where boq_version_id = :v and work_item_code = 'CONC-COL'"), {"v": v1})).scalar_one())
  ex = _ok(await client.get(f"{P}/boq-items/{item_id}/rate-resolution"))
  assert ex["current_source"] == "RATE_BOOK" and ex["would_resolve_to"]["unit_rate"] == "20000.00" and ex["stack"][0]["code"] == "CSR"
  _err(await client.get(f"{P}/boq-items/{uuid4()}/rate-resolution"), 404, "BOQ_ITEM_NOT_FOUND")

  engine = BOQEngineService(db_session)
  await engine.submit_for_review(organization.id, v1, actor_a.user_id)
  await db_session.refresh(await db_session.get(BOQVersion, v1))
  await engine.approve_version(organization.id, v1, actor_a.user_id)
  v2 = await _built(db_session, organization, actor_a, project_a, make_run, [("CONC-COL", "12", "m3", "G1"), ("STEEL-60", "100", "kg", "G2")])
  d = _ok(await client.get(f"{P}/boq-versions/{v1}/diff/{v2}"))
  assert d["summary"]["CHANGED"] == 1 and d["summary"]["UNCHANGED"] == 1 and [l["work_item_code"] for l in d["lines"]] == ["CONC-COL"]
  assert d["lines"][0]["quantity_delta"] == "2.0000" and d["summary"]["total_delta"] == "40000.00"
  full = _ok(await client.get(f"{P}/boq-versions/{v1}/diff/{v2}", params={"include_unchanged": True, "include_elements": False}))
  assert len(full["lines"]) == 2
  _err(await client.get(f"{P}/boq-versions/{v1}/diff/{v1}"), 422, "DIFF_SAME_VERSION")
  _err(await client.get(f"{P}/boq-versions/{v1}/diff/{uuid4()}"), 404, "BOQ_VERSION_NOT_FOUND")

  await engine.submit_for_review(organization.id, v2, actor_a.user_id)
  await db_session.refresh(await db_session.get(BOQVersion, v2))
  await engine.approve_version(organization.id, v2, actor_a.user_id)
  r = await client.get(f"{P}/boq-versions/{v2}/exports/REVISION_COMPARISON", params={"fmt": "xlsx"})
  assert r.status_code == 200 and "-vs-s" in r.headers["content-disposition"]
  assert {"Summary", "Changes"} <= set(load_workbook(BytesIO(r.content)).sheetnames)
  assert (await client.get(f"{P}/boq-versions/{v2}/exports/REVISION_COMPARISON", params={"fmt": "pdf"})).status_code == 422
  assert (await client.get(f"{P}/boq-versions/{v2}/exports/REVISION_COMPARISON", params={"fmt": "xlsx", "compare_snapshot_id": str(uuid4())})).status_code == 404

READ_ONLY = [PermissionKey.DRAWING_READ]

@pytest.mark.parametrize("method,path,body", [
  ("post", "/rate-books", {"code": "X", "name": "x"}),
  ("patch", "/rate-books/{id}", {"name": "x"}),
  ("delete", "/rate-books/{id}", None),
  ("post", "/rate-books/{id}/publish", None),
  ("post", "/rate-books/{id}/new-version", None),
  ("post", "/rate-books/{id}/archive", None),
  ("post", "/rate-books/{id}/items", {"work_item_code": "A", "unit": "m3", "rate": "1"}),
  ("post", "/rate-books/{id}/items/bulk", {"items": []}),
  ("patch", "/rate-items/{id}", {"rate": "1"}),
  ("delete", "/rate-items/{id}", None),
  ("post", "/rate-books/{id}/escalations", {"effective_from": "2026-01-01", "factor": "1.1"}),
  ("delete", "/rate-escalations/{id}", None),
  ("post", "/rate-books/{id}/analyses", {"code": "A", "work_item_code": "A", "description": "d", "unit": "m3"}),
  ("patch", "/rate-analyses/{id}", {"description": "x"}),
  ("delete", "/rate-analyses/{id}", None),
  ("post", "/rate-analyses/{id}/compute", None),
  ("post", "/rate-analyses/{id}/apply", None),
  ("post", "/projects/{id}/rate-overrides", {"work_item_code": "A", "unit": "m3", "rate": "1", "reason": "r"}),
  ("post", "/rate-overrides/{id}/revoke", {"reason": "r"}),
  ("post", "/boq-versions/{id}/price", None),
])
async def test_write_endpoints_need_permission(client, make_actor, organization, act_as, method, path, body):
  viewer = await make_actor(organization, READ_ONLY)
  act_as(viewer)
  r = await getattr(client, method)(f"{P}{path.format(id=uuid4())}", **({"json": body} if body is not None else {}))
  assert r.status_code == 403, r.text

@pytest.mark.parametrize("path", [
  "/rate-books", "/rate-books/{id}", "/rate-books/{id}/items", "/rate-books/{id}/escalations", "/rate-books/{id}/analyses",
  "/rate-analyses/{id}", "/rate-analyses/{id}/breakdown", "/projects/{id}/rate-overrides", "/boq-versions/{id}/pricing",
  "/boq-items/{id}/rate-resolution", "/boq-versions/{id}/diff/{id2}",
])
async def test_read_endpoints_need_drawing_read(client, make_actor, organization, act_as, path):
  nobody = await make_actor(organization, [])
  act_as(nobody)
  assert (await client.get(f"{P}{path.format(id=uuid4(), id2=uuid4())}")).status_code == 403

async def test_read_only_user_can_read_but_not_write(client, make_actor, organization, actor_a, act_as, make_book):
  await make_book(organization, "OURS", rates=[("A", "m3", 1)])
  viewer = await make_actor(organization, READ_ONLY)
  act_as(viewer)
  assert len(_ok(await client.get(f"{P}/rate-books"))) == 1

async def test_override_permission_is_boq_adjust(client, make_actor, organization, project_a, act_as):
  only_ratebook = await make_actor(organization, [PermissionKey.DRAWING_READ, PermissionKey.RATEBOOK_MANAGE])
  act_as(only_ratebook)
  r = await client.post(f"{P}/projects/{project_a.id}/rate-overrides", json={"work_item_code": "A", "unit": "m3", "rate": "1", "reason": "r"})
  assert r.status_code == 403
  adjuster = await make_actor(organization, [PermissionKey.DRAWING_READ, PermissionKey.BOQ_ADJUST])
  act_as(adjuster)
  assert (await client.post(f"{P}/projects/{project_a.id}/rate-overrides", json={"work_item_code": "A", "unit": "m3", "rate": "1", "reason": "r"})).status_code == 201

async def test_cross_tenant_access_is_404(client, make_actor, organization, other_organization, actor_a, actor_b, act_as, make_book, project_b, project_a):
  theirs = await make_book(other_organization, "THEIRS", rates=[("A", "m3", 1)])
  act_as(actor_a)
  assert _ok(await client.get(f"{P}/rate-books")) == []
  
  for method, path, body in (
    ("get", f"/rate-books/{theirs.id}", None), ("patch", f"/rate-books/{theirs.id}", {"name": "x"}),
    ("post", f"/rate-books/{theirs.id}/publish", None), ("get", f"/rate-books/{theirs.id}/items", None),
    ("post", f"/rate-books/{theirs.id}/items", {"work_item_code": "A", "unit": "m3", "rate": "1"}),
    ("post", f"/rate-books/{theirs.id}/new-version", None), ("delete", f"/rate-books/{theirs.id}", None),
    ("get", f"/rate-books/{theirs.id}/escalations", None), ("get", f"/rate-books/{theirs.id}/analyses", None),
    ("get", f"/projects/{project_b.id}/rate-overrides", None)):
    r = await getattr(client, method)(f"{P}{path}", **({"json": body} if body is not None else {}))
    assert r.status_code == 404, (method, path, r.status_code, r.text)
  _err(await client.post(f"{P}/projects/{project_b.id}/rate-overrides", json={"work_item_code": "A", "unit": "m3", "rate": "1", "reason": "r"}), 404, "PROJECT_NOT_FOUND")
  act_as(actor_b)
  assert [b["code"] for b in _ok(await client.get(f"{P}/rate-books"))] == ["THEIRS"]

async def test_system_books_are_visible_but_read_only_over_http(client, actor_a, act_as, make_book):
  act_as(actor_a)
  system = await make_book(None, "CSR", rates=[("CONC-COL", "m3", 21000)], escalations=[("ALL", __import__("datetime").date(2026, 1, 1), "1.1")])
  assert _ok(await client.get(f"{P}/rate-books/{system.id}"))["is_system"] is True
  assert len(_ok(await client.get(f"{P}/rate-books/{system.id}/items"))) == 1
  assert len(_ok(await client.get(f"{P}/rate-books/{system.id}/escalations"))) == 1
  for method, path, body in (("patch", "", {"name": "x"}), ("delete", "", None), ("post", "/publish", None), ("post", "/archive", None),
    ("post", "/items", {"work_item_code": "A", "unit": "m3", "rate": "1"}), ("post", "/escalations", {"effective_from": "2027-01-01", "factor": "1.1"})):
    r = await getattr(client, method)(f"{P}/rate-books/{system.id}{path}", **({"json": body} if body is not None else {}))
    assert r.status_code == 403, (method, path, r.status_code, r.text)
  item_id = _ok(await client.get(f"{P}/rate-books/{system.id}/items"))[0]["id"]
  
  assert (await client.patch(f"{P}/rate-items/{item_id}", json={"rate": "1"})).status_code == 403
  assert (await client.delete(f"{P}/rate-items/{item_id}")).status_code == 403
  copy = _ok(await client.post(f"{P}/rate-books", json={"code": "MY-CSR", "name": "Our copy", "copy_from_id": str(system.id)}), 201)
  assert copy["item_count"] == 1 and copy["is_system"] is False
  assert _ok(await client.patch(f"{P}/rate-books/{copy['id']}", json={"name": "Edited"}))["name"] == "Edited"