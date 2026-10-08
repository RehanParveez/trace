from __future__ import annotations
from datetime import date
from decimal import Decimal as D
from io import BytesIO
import pytest
from openpyxl import load_workbook
from app.core.exceptions import TraceException
from app.modules.drawings_boq.boq_service import BOQEngineService
from app.modules.drawings_boq.models import BOQSnapshot, BOQVersion
from app.modules.drawings_boq.pricing.diff_service import BOQDiffService
from app.modules.drawings_boq.pricing.pricing_service import PricingService
from app.modules.drawings_boq.pricing.rate_service import RateBookService
from app.modules.drawings_boq.pricing.schemas import EscalationCreateRequest, OverrideCreateRequest, PriceVersionRequest, RateBookCreateRequest, RateItemInput, RateItemUpdateRequest
from app.modules.drawings_boq.repository import BOQItemRepository

async def _fresh(db, *version_ids):
  for vid in version_ids:
    await db.refresh(await db.get(BOQVersion, vid))

def q2(x):
  return D(x).quantize(D("0.01"), rounding="ROUND_HALF_UP")

async def _sheet_rows(data: bytes, name: str):
  wb = load_workbook(BytesIO(data))
  return wb[name], [[c.value for c in row] for row in wb[name].iter_rows()]

async def test_full_pricing_flow_from_rate_book_to_revision_comparison(db_session, organization, actor_a, project_a, make_run):
  org, uid, pid = organization.id, actor_a.user_id, project_a.id
  rates, engine, pricing, diffs = RateBookService(db_session), BOQEngineService(db_session), PricingService(db_session), BOQDiffService(db_session)

  book = await rates.create_book(org, uid, RateBookCreateRequest(code="OURS", name="Our rates", edition="2026"))
  for code, unit, rate in (("CONC-COL", "m3", "20000"), ("STEEL-60", "kg", "280"), ("BRICK", "nos", "28")):
    await rates.add_item(org, uid, book.id, RateItemInput(work_item_code=code, unit=unit, rate=D(rate)))
  await rates.add_escalation(org, uid, book.id, EscalationCreateRequest(effective_from=date(2020, 1, 1), factor=D("1.05")))
  v1_book = await rates.publish_book(org, uid, book.id)
  assert v1_book.status == "ACTIVE"

  run1 = await make_run(organization, project_a, [("CONC-COL", "10", "m3", "G1"), ("STEEL-60", "1000", "kg", "G2")])
  built1 = await engine.build_from_run(org, run1.id, uid)
  v1_id = built1["boq_version_id"]
  items1 = {i.work_item_code: i for i in await BOQItemRepository(db_session).list_by_version(v1_id, org)}
  assert items1["CONC-COL"].unit_rate == D("21000.00") and items1["CONC-COL"].base_rate == D("20000.00")
  assert items1["STEEL-60"].unit_rate == D("294.00")
  assert built1["open_issues"] == 0

  await rates.create_override(org, uid, pid, OverrideCreateRequest(work_item_code="STEEL-60", unit="kg", rate=D("300"), reason="Quote from supplier"))
  out = await pricing.price_version(org, v1_id, uid, PriceVersionRequest())
  assert out["by_source"] == {"RATE_BOOK": 1, "PROJECT_OVERRIDE": 1} and out["unpriced"] == 0
  expected_total = q2(D("10") * D("21000")) + q2(D("1000") * D("300"))
  assert D(out["total"]) == expected_total == D("510000.00")

  await engine.submit_for_review(org, v1_id, uid)
  await _fresh(db_session, v1_id)
  await engine.approve_version(org, v1_id, uid)
  await _fresh(db_session, v1_id)
  v1 = await db_session.get(BOQVersion, v1_id)
  assert v1.lifecycle == "APPROVED" and v1.snapshot_id is not None
  snap1 = await db_session.get(BOQSnapshot, v1.snapshot_id)
  assert D(snap1.totals["grand"]) == expected_total
  rows = await engine.snapshot_items(org, snap1.id)
  by_code = {r.work_item_code: r for r in rows}
  assert by_code["CONC-COL"].base_rate == D("20000.00") and by_code["CONC-COL"].escalation_factor == D("1.0500")
  assert by_code["STEEL-60"].rate_resolution["source"] == "PROJECT_OVERRIDE"
  await engine.issue_version(org, v1_id, uid)
  await _fresh(db_session, v1_id)
  with pytest.raises(TraceException) as exc:                   
    await pricing.price_version(org, v1_id, uid, PriceVersionRequest())
  assert exc.value.code == "BOQ_IMMUTABLE"

  draft = await rates.new_version(org, uid, v1_book.id)
  conc = next(i for i in (await rates.list_items(org, draft.id, q="conc", trade=None, only_active=True, after=None, limit=5))[0])
  await rates.update_item(org, uid, conc.id, RateItemUpdateRequest(rate=D("22000")))
  v2_book = await rates.publish_book(org, uid, draft.id)
  assert v2_book.supersedes_rate_book_id == v1_book.id

  run2 = await make_run(organization, project_a, [("CONC-COL", "5", "m3", "G1"), ("CONC-COL", "7", "m3", "G3"), ("BRICK", "500", "nos", "G4")])
  built2 = await engine.build_from_run(org, run2.id, uid)
  v2_id = built2["boq_version_id"]
  assert v2_id != v1_id
  items2 = {i.work_item_code: i for i in await BOQItemRepository(db_session).list_by_version(v2_id, org)}
  assert set(items2) == {"CONC-COL", "BRICK"}
  assert items2["CONC-COL"].quantity == D("12.0000") and items2["CONC-COL"].unit_rate == D("23100.00")     
  assert items2["BRICK"].unit_rate == D("29.40")
  assert items2["CONC-COL"].rate_book_id == v2_book.id

  diff = await diffs.diff_versions(org, v1_id, v2_id)
  status = {l["work_item_code"]: l for l in diff["lines"]}
  assert {k: v["status"] for k, v in status.items()} == {"CONC-COL": "CHANGED", "STEEL-60": "REMOVED", "BRICK": "ADDED"}
  assert status["CONC-COL"]["quantity_delta"] == D("2.0000") and status["CONC-COL"]["rate_delta"] == D("2100.00")
  assert [e["ifc_global_id"] for e in status["CONC-COL"]["elements_added"]] == ["G3"]
  assert status["STEEL-60"]["amount_delta"] == D("-300000.00")
  s = diff["summary"]
  assert (s["ADDED"], s["REMOVED"], s["CHANGED"]) == (1, 1, 1)
  assert s["total_delta"] == s["total_b"] - s["total_a"]

  await engine.submit_for_review(org, v2_id, uid)
  await _fresh(db_session, v2_id)
  await engine.approve_version(org, v2_id, uid)
  await _fresh(db_session, v2_id, v1_id)
  assert (await db_session.get(BOQVersion, v1_id)).lifecycle == "SUPERSEDED"
  data, media, name = await engine.export_snapshot(org, v2_id, "REVISION_COMPARISON", "xlsx", uid)
  assert name.endswith(".xlsx") and "-vs-s" in name and media.endswith("spreadsheetml.sheet")
  sheet, grid = await _sheet_rows(data, "Changes")
  flat = " ".join(str(c) for r in grid for c in r if c is not None)
  for word in ("CONC-COL", "STEEL-60", "BRICK", "ADDED", "REMOVED", "CHANGED"):
    assert word in flat
  contract, _, _ = await engine.export_snapshot(org, v2_id, "CONTRACT_BOQ", "xlsx", uid)
  _, cgrid = await _sheet_rows(contract, load_workbook(BytesIO(contract)).sheetnames[0])
  assert any("Base rate" in str(c) for r in cgrid for c in r) and any("Escalation" in str(c) for r in cgrid for c in r)

async def test_comparison_export_errors(db_session, organization, actor_a, project_a, make_run, make_book):
  org, uid = organization.id, actor_a.user_id
  await make_book(None, "CSR", rates=[("CONC-COL", "m3", 20000)])
  engine = BOQEngineService(db_session)
  run = await make_run(organization, project_a, [("CONC-COL", "10", "m3", "G1")])
  vid = (await engine.build_from_run(org, run.id, uid))["boq_version_id"]
  await engine.submit_for_review(org, vid, uid)
  await _fresh(db_session, vid)
  await engine.approve_version(org, vid, uid)
  await _fresh(db_session, vid)
  with pytest.raises(TraceException) as exc:
    await engine.export_snapshot(org, vid, "REVISION_COMPARISON", "xlsx", uid)
  assert exc.value.code == "NO_BASE_SNAPSHOT"
  snap_id = (await db_session.get(BOQVersion, vid)).snapshot_id
  with pytest.raises(TraceException) as exc:
    await engine.export_snapshot(org, vid, "REVISION_COMPARISON", "xlsx", uid, compare_snapshot_id=snap_id)
  assert exc.value.code == "DIFF_SAME_VERSION"
  with pytest.raises(TraceException) as exc:
    await engine.export_snapshot(org, vid, "REVISION_COMPARISON", "pdf", uid)
  assert exc.value.code == "EXPORT_FORMAT_UNSUPPORTED"

async def test_diff_guards(db_session, organization, other_organization, actor_a, actor_b, project_a, project_b, make_run, make_book):
  org, uid = organization.id, actor_a.user_id
  await make_book(None, "CSR", rates=[("CONC-COL", "m3", 20000)])
  engine, diffs = BOQEngineService(db_session), BOQDiffService(db_session)
  r1 = await make_run(organization, project_a, [("CONC-COL", "10", "m3", "G1")])
  a_id = (await engine.build_from_run(org, r1.id, uid))["boq_version_id"]
  rb = await make_run(other_organization, project_b, [("CONC-COL", "10", "m3", "G1")])
  b_id = (await engine.build_from_run(other_organization.id, rb.id, actor_b.user_id))["boq_version_id"]
  for coro, code in ((diffs.diff_versions(org, a_id, a_id), "DIFF_SAME_VERSION"),
    (diffs.diff_versions(org, a_id, b_id), "BOQ_VERSION_NOT_FOUND")):
    with pytest.raises(TraceException) as exc:
      await coro
    assert exc.value.code == code

async def test_diff_unchanged_lines_hidden_unless_requested(db_session, organization, actor_a, project_a, make_run, make_book):
  org, uid = organization.id, actor_a.user_id
  await make_book(None, "CSR", rates=[("CONC-COL", "m3", 20000), ("STEEL-60", "kg", 280)])
  engine, diffs = BOQEngineService(db_session), BOQDiffService(db_session)
  r1 = await make_run(organization, project_a, [("CONC-COL", "10", "m3", "G1"), ("STEEL-60", "100", "kg", "G2")])
  v1 = (await engine.build_from_run(org, r1.id, uid))["boq_version_id"]
  await engine.submit_for_review(org, v1, uid)
  await _fresh(db_session, v1)
  await engine.approve_version(org, v1, uid)
  await _fresh(db_session, v1)
  r2 = await make_run(organization, project_a, [("CONC-COL", "10", "m3", "G1"), ("STEEL-60", "150", "kg", "G2")])
  v2 = (await engine.build_from_run(org, r2.id, uid))["boq_version_id"]
  hidden = await diffs.diff_versions(org, v1, v2)
  assert [l["work_item_code"] for l in hidden["lines"]] == ["STEEL-60"]
  shown = await diffs.diff_versions(org, v1, v2, include_unchanged=True)
  assert {l["work_item_code"]: l["status"] for l in shown["lines"]} == {"CONC-COL": "UNCHANGED", "STEEL-60": "CHANGED"}